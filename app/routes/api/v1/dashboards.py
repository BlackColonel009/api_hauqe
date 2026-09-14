"""
Routes API — Pilotage / Tableaux de bord / Baromètre / Public.

ROUTES INTERNES
---------------
- `/api/v1/dashboards/operational`
- `/api/v1/dashboards/tactical`
- `/api/v1/dashboards/strategic`
- `/api/v1/dashboards/annual`
- `/api/v1/barometer`

ROUTE PUBLIQUE
--------------
- `/api/v1/public/indicators`

La route publique n'exige pas de token, mais le service refuse toute réponse
si aucune règle + publication institutionnelle valide ne l'autorise.
"""

from __future__ import annotations

from datetime import date, datetime, timezone
import io
from uuid import UUID

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.session import get_db
from app.permissions.auth import require_permission
from app.schemas.dashboard import (
    AnnualDashboardResponse,
    BarometerResponse,
    DashboardFiltersResponse,
    IndicatorDefinitionsResponse,
    OperationalDashboardResponse,
    PublicIndicatorsResponse,
    StrategicDashboardResponse,
    TacticalDashboardResponse,
)
from app.services.auth_service import AuthContext
from app.services.dashboard_service import DashboardService


dashboard_router = APIRouter(
    prefix="/dashboards",
    tags=["Pilotage / Tableaux de bord"],
)

barometer_router = APIRouter(
    prefix="/barometer",
    tags=["Baromètre national"],
)

public_dashboard_router = APIRouter(
    prefix="/public",
    tags=["Données publiques agrégées"],
)


def _write_dashboard_section(
    sheet,
    *,
    row: int,
    title: str,
    headers: list[str],
    rows: list[list[object]],
) -> int:
    """Ajoute un tableau lisible dans l'export Excel opérationnel."""
    dark_green = "165D45"
    header_green = "23795A"
    light_green = "EAF5EF"
    thin = Side(style="thin", color="D9E5DF")

    sheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=len(headers),
    )
    title_cell = sheet.cell(row=row, column=1, value=title)
    title_cell.fill = PatternFill("solid", fgColor=dark_green)
    title_cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 22

    header_row = row + 1
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor=header_green)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    sheet.row_dimensions[header_row].height = 20

    for index, values in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=index, column=column, value=value)
            if isinstance(value, (date, datetime)):
                cell.number_format = "dd/mm/yyyy"
            cell.font = Font(name="Arial", size=10, color="1F352B")
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin)
            if (index - header_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light_green)
        sheet.row_dimensions[index].height = 34

    return header_row + max(1, len(rows)) + 2


def _operational_dashboard_xlsx(dashboard) -> bytes:
    """Construit un export Excel destiné à la lecture opérationnelle."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tableau de bord"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:D1")
    title = sheet["A1"]
    title.value = "HAUQE Certif — Tableau de bord opérationnel"
    title.font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="103F30")
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    generated_at = dashboard.generated_at.astimezone(timezone.utc)
    metadata = [
        ("Période", dashboard.period.label),
        ("Généré le", generated_at.replace(tzinfo=None)),
    ]
    for row, (label, value) in enumerate(metadata, start=3):
        sheet.cell(row=row, column=1, value=label).font = Font(
            name="Arial", size=10, bold=True, color="355347"
        )
        value_cell = sheet.cell(row=row, column=2, value=value)
        value_cell.font = Font(name="Arial", size=10, color="1F352B")
        if isinstance(value, datetime):
            value_cell.number_format = "dd/mm/yyyy hh:mm"

    row = 6
    row = _write_dashboard_section(
        sheet,
        row=row,
        title="INDICATEURS",
        headers=["Indicateur", "Valeur", "Unité"],
        rows=[
            [item.label, item.value, item.unit or ""]
            for item in dashboard.kpis
        ],
    )
    row = _write_dashboard_section(
        sheet,
        row=row,
        title="CERTIFICATIONS À ÉCHÉANCE",
        headers=["Entreprise", "Certification", "Expiration", "Jours restants"],
        rows=[
            [
                item.enterprise_name,
                item.certification_code or "Non renseignée",
                item.expiration_date,
                item.days_remaining,
            ]
            for item in dashboard.expiring_certifications
        ] or [["Aucune certification à échéance", "", "", ""]],
    )
    _write_dashboard_section(
        sheet,
        row=row,
        title="ACTIONS PRIORITAIRES",
        headers=["Action prioritaire", "Contexte", "Type", "Échéance"],
        rows=[
            [
                item.title,
                " · ".join(
                    value
                    for value in (
                        item.resource_label,
                        item.resource_subtitle,
                    )
                    if value
                ),
                "Alerte" if item.type == "ALERTE" else "Échéance",
                item.due_date,
            ]
            for item in dashboard.priority_actions
        ] or [["Aucune action prioritaire", "", "", ""]],
    )

    widths = {"A": 40, "B": 43, "C": 20, "D": 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ============================================================
# MÉTADONNÉES PARTAGÉES
# ============================================================

@dashboard_router.get(
    "/filters",
    response_model=DashboardFiltersResponse,
)
async def dashboard_filters(
    db: AsyncSession = Depends(get_db),
    actor: AuthContext = Depends(
        require_permission("DASHBOARDS.LIRE_REFERENTIELS")
    ),
):
    """
    Alimente les filtres région/secteur/norme/organisme du frontend.
    """
    return await DashboardService.filters(db)


@dashboard_router.get(
    "/indicator-definitions",
    response_model=IndicatorDefinitionsResponse,
)
async def indicator_definitions(
    actor: AuthContext = Depends(
        require_permission("DASHBOARDS.LIRE_REFERENTIELS")
    ),
):
    """
    Définitions fonctionnelles affichables dans les infobulles du frontend.
    """
    return DashboardService.definitions()


# ============================================================
# OPÉRATIONNEL
# ============================================================

@dashboard_router.get(
    "/operational/export",
)
async def operational_dashboard_export(
    days: int = Query(default=7, ge=1, le=90),
    zone_id: UUID | None = Query(default=None),
    sector: str | None = Query(default=None, max_length=255),
    norm_id: UUID | None = Query(default=None),
    organisme_id: UUID | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
    actor: AuthContext = Depends(
        require_permission("DASHBOARDS.OPERATIONNEL")
    ),
):
    """Export Excel mis en forme du même snapshot opérationnel."""
    dashboard = await DashboardService.operational(
        db,
        days=days,
        zone_id=zone_id,
        sector=sector,
        norm_id=norm_id,
        organisme_id=organisme_id,
    )

    return Response(
        content=_operational_dashboard_xlsx(dashboard),
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                "attachment; "
                'filename="hauqe-dashboard-operationnel.xlsx"'
            )
        },
    )


@dashboard_router.get(
    "/operational",
    response_model=OperationalDashboardResponse,
)
async def operational_dashboard(
    days: int = Query(default=7, ge=1, le=90),
    zone_id: UUID | None = Query(default=None),
    sector: str | None = Query(default=None, max_length=255),
    norm_id: UUID | None = Query(default=None),
    organisme_id: UUID | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
    actor: AuthContext = Depends(
        require_permission("DASHBOARDS.OPERATIONNEL")
    ),
):
    return await DashboardService.operational(
        db,
        days=days,
        zone_id=zone_id,
        sector=sector,
        norm_id=norm_id,
        organisme_id=organisme_id,
    )


# ============================================================
# TACTIQUE MENSUEL
# ============================================================

@dashboard_router.get(
    "/tactical",
    response_model=TacticalDashboardResponse,
)
async def tactical_dashboard(
    year: int = Query(ge=2000, le=2100),
    month: int = Query(ge=1, le=12),
    db: AsyncSession = Depends(get_db),
    actor: AuthContext = Depends(
        require_permission("DASHBOARDS.TACTIQUE")
    ),
):
    return await DashboardService.tactical(
        db,
        year=year,
        month=month,
    )


# ============================================================
# STRATÉGIQUE TRIMESTRIEL
# ============================================================

@dashboard_router.get(
    "/strategic",
    response_model=StrategicDashboardResponse,
)
async def strategic_dashboard(
    year: int = Query(ge=2000, le=2100),
    quarter: int = Query(ge=1, le=4),
    db: AsyncSession = Depends(get_db),
    actor: AuthContext = Depends(
        require_permission("DASHBOARDS.STRATEGIQUE")
    ),
):
    return await DashboardService.strategic(
        db,
        year=year,
        quarter=quarter,
    )


# ============================================================
# ANNUEL
# ============================================================

@dashboard_router.get(
    "/annual",
    response_model=AnnualDashboardResponse,
)
async def annual_dashboard(
    year: int = Query(ge=2000, le=2100),
    db: AsyncSession = Depends(get_db),
    actor: AuthContext = Depends(
        require_permission("DASHBOARDS.ANNUEL")
    ),
):
    return await DashboardService.annual(
        db,
        year=year,
    )


# ============================================================
# BAROMÈTRE NATIONAL
# ============================================================

@barometer_router.get(
    "",
    response_model=BarometerResponse,
)
async def national_barometer(
    start_date: date | None = Query(default=None),
    end_date: date | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
    actor: AuthContext = Depends(
        require_permission("BAROMETRE.LIRE")
    ),
):
    today = date.today()
    start = start_date or date(today.year, 1, 1)
    end = end_date or today

    return await DashboardService.barometer(
        db,
        start_date=start,
        end_date=end,
    )


# ============================================================
# TABLEAU PUBLIC AGRÉGÉ
# ============================================================

@public_dashboard_router.get(
    "/indicators",
    response_model=PublicIndicatorsResponse,
)
async def public_indicators(
    response: Response,
    db: AsyncSession = Depends(get_db),
):
    """
    Endpoint volontairement anonyme.

    La couche service vérifie néanmoins :
    - règle publiée ;
    - allowlist d'indicateurs ;
    - période publiée ;
    - publication institutionnelle PUBLIEE.

    Cache court uniquement : une nouvelle publication doit devenir visible
    rapidement.
    """
    response.headers["Cache-Control"] = "public, max-age=300"
    response.headers["X-Content-Type-Options"] = "nosniff"
    return await DashboardService.public_indicators(db)
