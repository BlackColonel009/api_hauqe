"""Bilans BNEC périodiques fondés sur les tableaux de bord serveur."""

from __future__ import annotations

import csv
import re
from datetime import date, datetime, timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException, Request
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.document import Document
from app.models.rapport_genere import RapportGenere
from app.repositories.governance_repository import GovernanceRepository
from app.repositories.dashboard_repository import DashboardRepository
from app.services.dashboard_service import DashboardService
from app.services.document_service import storage_root
from app.services.governance_service import GovernanceService
from app.audit.service import write_audit_event


TEMPLATES = {
    "MENSUEL": ("Bilan mensuel BNEC", "month", DashboardService.tactical),
    "TRIMESTRIEL": ("Bilan trimestriel BNEC", "quarter", DashboardService.strategic),
    "ANNUEL": ("Rapport annuel BNEC", "year", DashboardService.annual),
}
FORMATS = {"PDF": ("pdf", "application/pdf"), "XLSX": ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), "CSV": ("csv", "text/csv; charset=utf-8")}
METHODOLOGY = (
    "Les activités sont calculées sur la période choisie. Les indicateurs de stock "
    "et les alertes actives reflètent l'état du registre au moment de la génération. "
    "Si la période n'est pas terminée, le bilan est provisoire."
)
LOGO_PATH = Path(__file__).resolve().parents[1] / "static" / "logo.jpg"
METRIC_LABELS = {
    "submitted_forms": "Fiches de collecte soumises", "opened": "Dossiers de vérification ouverts",
    "closed": "Dossiers de vérification clôturés", "finalized": "Contrôles FUCCS finalisés",
    "average_rate": "Taux moyen des contrôles FUCCS", "decisions": "Décisions de validation",
    "completed": "Intégrations BNEC terminées", "new_certifications": "Certifications enregistrées",
    "alerts_created": "Alertes créées", "alerts_resolved": "Alertes résolues",
    "renewal_decisions": "Décisions de renouvellement", "reviews_validated": "Revues qualité validées",
    "open_action_plans": "Plans d'action encore ouverts", "open_action_plans_at_generation": "Plans d'action encore ouverts",
    "incidents_declared": "Incidents déclarés", "quality_reviews_created": "Revues qualité créées",
    "backup_failures": "Échecs de sauvegarde",
}
METRIC_EXPLANATIONS = {
    "submitted_forms": "Fiches transmises pour traitement pendant la période.",
    "opened": "Vérifications commencées pendant la période.",
    "closed": "Vérifications terminées pendant la période.",
    "finalized": "Contrôles FUCCS terminés pendant la période.",
    "average_rate": "Moyenne des taux des contrôles FUCCS finalisés pendant la période.",
    "completed": "Dossiers intégrés à la BNEC pendant la période.",
    "new_certifications": "Certifications créées dans le registre pendant la période.",
    "alerts_created": "Alertes enregistrées pendant la période.",
    "alerts_resolved": "Alertes résolues pendant la période.",
    "renewal_decisions": "Décisions de renouvellement prises pendant la période.",
    "reviews_validated": "Revues qualité validées pendant la période.",
    "open_action_plans": "Plans non clôturés au moment de la génération.",
    "open_action_plans_at_generation": "Plans non clôturés au moment de la génération.",
    "incidents_declared": "Incidents déclarés pendant la période.",
    "quality_reviews_created": "Revues qualité créées pendant la période.",
    "backup_failures": "Sauvegardes en échec pendant la période.",
}
SECTION_LABELS = {
    "collection": "Collecte", "verification": "Vérification documentaire",
    "fuccs": "Contrôle FUCCS", "validation": "Validation N1/N2",
    "integration": "Intégration BNEC", "watch": "Veille et alertes",
    "quality": "Qualité des données", "governance": "Gouvernance", "continuity": "Continuité",
}
DISTRIBUTION_LABELS = {
    "certification_statuses": "Statuts actuels des certifications",
    "sncc_classes": "Classes SNCC actuelles", "sncc_risks": "Niveaux de risque SNCC actuels",
    "by_sector": "Secteurs du registre actuel", "by_norm": "Normes du registre actuel",
    "by_certification_body": "Organismes certificateurs du registre actuel",
}
KPI_EXPLANATIONS = {
    "submitted_collection_forms": "Fiches de collecte soumises pendant la période.",
    "new_certifications": "Certifications créées dans le registre pendant la période.",
    "fuccs_finalized": "Contrôles FUCCS finalisés pendant la période.",
    "integrations_completed": "Dossiers intégrés à la BNEC pendant la période.",
    "alerts_created": "Alertes enregistrées pendant la période.",
    "enterprises_count": "Entreprises présentes dans le registre au moment de la génération.",
    "certifications_count": "Certifications présentes dans le registre au moment de la génération.",
    "active_certifications_count": "Certifications au statut actif au moment de la génération.",
    "validated_infc_latest_count": "Certifications disposant d'un dernier INFC validé.",
    "quarter_infc_average": "Moyenne des INFC calculés pour le trimestre choisi.",
    "critical_alerts": "Alertes critiques encore actives au moment de la génération.",
    "overdue_deadlines": "Échéances non terminées dont la date est dépassée.",
    "renewal_decisions": "Décisions de renouvellement prises pendant la période.",
    "latest_national_infc_average": "Moyenne des derniers INFC validés disponibles au moment de la génération.",
}


def validate_period(kind: str, year: int, month: int | None, quarter: int | None) -> tuple[date, date, str]:
    if kind not in TEMPLATES or year < 2000 or year > 2100:
        raise HTTPException(422, "Type ou année de rapport invalide.")
    if kind == "MENSUEL":
        if month is None or not 1 <= month <= 12:
            raise HTTPException(422, "Choisissez un mois valide.")
        import calendar
        return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1]), f"{month:02d}/{year}"
    if kind == "TRIMESTRIEL":
        if quarter is None or not 1 <= quarter <= 4:
            raise HTTPException(422, "Choisissez un trimestre valide.")
        import calendar
        first_month = (quarter - 1) * 3 + 1
        last_month = first_month + 2
        return date(year, first_month, 1), date(year, last_month, calendar.monthrange(year, last_month)[1]), f"T{quarter} {year}"
    return date(year, 1, 1), date(year, 12, 31), str(year)


def scalar(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    if isinstance(value, (float, Decimal)):
        return f"{value:,.2f}".replace(",", " ").replace(".", ",")
    if isinstance(value, str) and re.fullmatch(r"-?\d+\.\d+", value):
        return f"{Decimal(value):,.2f}".replace(",", " ").replace(".", ",")
    return str(value)


def readable_rows(data: dict, kind: str, period: str, *, as_of: date | None = None) -> list[tuple[str, str]]:
    """Présente les agrégats serveur en langage métier, sans champs ni UUID techniques."""
    rows: list[tuple[str, str]] = []
    kpis = data.get("kpis") or []
    by_key = {item.get("key"): item.get("value") for item in kpis}
    if kind == "MENSUEL":
        intro = (f"En {period}, {scalar(by_key.get('submitted_collection_forms'))} fiche(s) de collecte "
                 f"ont été soumises et {scalar(by_key.get('integrations_completed'))} dossier(s) ont été intégrés à la BNEC.")
    elif kind == "TRIMESTRIEL":
        intro = (f"Pour {period}, le rapport présente l'INFC moyen du trimestre et "
                 "l'état actuel du registre, des risques et des répartitions.")
    else:
        intro = (f"Pour {period}, {scalar(by_key.get('new_certifications'))} certification(s) ont été "
                 f"enregistrées et {scalar(by_key.get('integrations_completed'))} dossier(s) intégrés à la BNEC.")
        if as_of and as_of.year == int(period) and as_of < date(int(period), 12, 31):
            intro = f"Bilan provisoire arrêté au {as_of.strftime('%d/%m/%Y')}. " + intro
    rows.append(("En bref", intro))
    if kind == "ANNUEL" and as_of and as_of.year == int(period) and as_of < date(int(period), 12, 31):
        rows.append(("Période observée", f"Du 01/01/{period} au {as_of.strftime('%d/%m/%Y')}. Les mois à venir ne sont pas comptés comme des résultats nuls."))
    for item in kpis:
        if item.get("value") is None:
            continue
        value = scalar(item["value"])
        if item.get("unit"):
            value += f" {item['unit']}"
        if item.get("previous_value") is not None:
            comparison = {"MENSUEL": "mois précédent", "TRIMESTRIEL": "trimestre précédent", "ANNUEL": "année précédente"}[kind]
            value += f" ; {comparison} : {scalar(item['previous_value'])}"
            if item.get("delta") is not None:
                delta_value = item["delta"]
                value += f" ; évolution : {'+' if float(delta_value) > 0 else ''}{scalar(delta_value)}"
        explanation = item.get("definition") or KPI_EXPLANATIONS.get(item.get("key"))
        if explanation:
            value += f". {explanation}"
        rows.append((f"Chiffres à retenir — {item['label']}", value))

    for section_key, title in SECTION_LABELS.items():
        section = data.get(section_key)
        if not isinstance(section, dict):
            continue
        for metric_key, metric_value in section.items():
            metric_label = METRIC_LABELS.get(metric_key, metric_key.replace("_", " ").capitalize())
            if isinstance(metric_value, dict):
                for decision, count in metric_value.items():
                    decision_label = str(decision).replace("_", " ").capitalize()
                    rows.append((f"{title} — {metric_label} : {decision_label}", scalar(count)))
            else:
                value = scalar(metric_value)
                if metric_key == "average_rate" and metric_value is not None:
                    value += " %"
                explanation = METRIC_EXPLANATIONS.get(metric_key)
                rows.append((f"{title} — {metric_label}", f"{value}. {explanation}" if explanation else value))

    for key, title in DISTRIBUTION_LABELS.items():
        for item in data.get(key) or []:
            name = item.get("label") or item.get("key") or "Non renseigné"
            if key == "certification_statuses" and str(name).casefold() == "a verifier":
                name = "À vérifier"
            value = f"{scalar(item.get('value'))} certification(s)"
            if item.get("percentage") is not None:
                value += f" ; part dans cette répartition : {scalar(item['percentage'])} %"
            rows.append((f"{title} — {name}", value))

    for item in data.get("by_region") or []:
        value = (f"{scalar(item.get('certifications'))} certification(s), dont "
                 f"{scalar(item.get('active_certifications'))} active(s) ; "
                 f"{scalar(item.get('enterprises'))} entreprise(s)")
        if item.get("average_infc") is not None:
            value += f" ; INFC moyen : {scalar(item['average_infc'])}"
        rows.append((f"Répartition régionale actuelle — {item.get('zone_name') or 'Zone non renseignée'}", value))

    for key, title in (
        ("infc_series", "Évolution de l'INFC"),
        ("quarterly_certifications", "Certifications enregistrées"),
        ("quarterly_infc", "INFC moyen"),
    ):
        for point in data.get(key) or []:
            point_period = point.get("period") or "Période non renseignée"
            point_end = None
            if kind == "ANNUEL" and as_of and re.fullmatch(r"\d{4}-T[1-4]", point_period):
                quarter_start, point_end, _ = validate_period("TRIMESTRIEL", int(point_period[:4]), None, int(point_period[-1]))
                if quarter_start > as_of:
                    continue
            point_value = point.get("value")
            if point_value is None and key in {"quarterly_infc", "infc_series"}:
                value = "Aucun résultat INFC validé pour ce trimestre."
            else:
                value = scalar(point_value)
            if point_end and as_of and as_of < point_end:
                value += f" (trimestre en cours, données au {as_of.strftime('%d/%m/%Y')})"
            rows.append((f"{title} — {point_period}", value))

    synthesis = data.get("synthesis") or {}
    for key, title in (
        ("findings", "Constat"), ("major_risks", "Risque à surveiller"),
        ("priority_recommendations", "Action recommandée"),
    ):
        for index, message in enumerate(synthesis.get(key) or [], 1):
            rows.append((f"Analyse — {title} {index}", str(message)))

    rows.extend([
        ("Repère de lecture — INFC", "Indice utilisé pour apprécier la fiabilité d'une certification ; sa valeur provient des résultats validés dans le système."),
        ("Repère de lecture — FUCCS", "Grille de contrôle appliquée au dossier après la vérification documentaire."),
        ("Repère de lecture — SNCC", "Classement des certifications selon les règles publiées dans le système."),
    ])
    return rows


def render_csv(title: str, period: str, rows: list[tuple[str, str]]) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream, delimiter=";")
    writer.writerow(["HAUQE - BNEC", title])
    writer.writerow(["Période", period])
    writer.writerow(["Méthode", METHODOLOGY])
    writer.writerow(["Indicateur", "Valeur"])
    writer.writerows(rows)
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def render_xlsx(title: str, period: str, rows: list[tuple[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bilan BNEC"
    sheet.append(["HAUQE - BNEC", title])
    sheet.append(["Période", period])
    sheet.append(["Méthode", METHODOLOGY])
    sheet.append(["Indicateur", "Valeur"])
    for row in rows:
        sheet.append(list(row))
    sheet.column_dimensions["A"].width = 72
    sheet.column_dimensions["B"].width = 32
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:B{sheet.max_row}"
    border = Border(bottom=Side(style="hair", color="D9D9D9"))
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor="E4E4E4")
        cell.font = Font(bold=True, color="000000")
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    sheet["A1"].font = Font(bold=True, size=14, color="000000")
    sheet["B1"].font = Font(bold=True, size=14, color="000000")
    if LOGO_PATH.is_file():
        logo = WorksheetImage(str(LOGO_PATH))
        logo.width = 52
        logo.height = 52
        sheet.add_image(logo, "D1")
        sheet.row_dimensions[1].height = 42
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_pdf(title: str, period: str, rows: list[tuple[str, str]]) -> bytes:
    output = BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=(595.28, 841.89), rightMargin=42, leftMargin=42, topMargin=48, bottomMargin=48)
    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.black
    styles["Heading1"].textColor = colors.black
    styles["Normal"].textColor = colors.black
    brand = Paragraph("HAUQE - BNEC", styles["Title"])
    if LOGO_PATH.is_file():
        identity = Table([[Image(str(LOGO_PATH), width=52, height=52), brand]], colWidths=[64, 447])
        identity.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story = [identity]
    else:
        story = [brand]
    story.extend([Spacer(1, 10), Paragraph(title, styles["Heading1"]), Paragraph(f"Période : {period}", styles["Normal"]), Spacer(1, 8)])
    if rows and rows[0][0] == "En bref":
        summary = rows[0][1].replace("&", "&amp;").replace("<", "&lt;")
        story.extend([Paragraph(f"<b>En bref :</b> {summary}", styles["Normal"]), Spacer(1, 8)])
    story.extend([Paragraph(METHODOLOGY, styles["Normal"]), Spacer(1, 18)])
    # Paragraphs wrap long indicator labels and values rather than clipping them.
    table_rows = [["Indicateur", "Valeur"]]
    for key, value in (rows[1:] if rows and rows[0][0] == "En bref" else rows):
        table_rows.append([Paragraph(key.replace("&", "&amp;").replace("<", "&lt;"), styles["Normal"]), Paragraph(value.replace("&", "&amp;").replace("<", "&lt;"), styles["Normal"])])
    table = Table(table_rows, colWidths=[355, 156], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E4E4E4")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    pdf.build(story)
    return output.getvalue()


class BnecReportService:
    @staticmethod
    async def build(db: AsyncSession, *, kind: str, year: int, month: int | None, quarter: int | None) -> dict:
        start, end, period_label = validate_period(kind, year, month, quarter)
        today = date.today()
        if start > today:
            raise HTTPException(422, "La période choisie n'a pas encore commencé.")
        as_of = min(today, end)
        if kind == "MENSUEL":
            dashboard = await DashboardService.tactical(db, year=year, month=month)
        elif kind == "TRIMESTRIEL":
            dashboard = await DashboardService.strategic(db, year=year, quarter=quarter)
        else:
            dashboard = await DashboardService.annual(db, year=year)
        data = dashboard.model_dump(mode="json")
        if kind == "ANNUEL":
            # Le tableau de bord emploie 0 comme valeur de remplacement pour
            # un trimestre sans INFC. La source permet de distinguer ce cas
            # d'un véritable score nul dans le rapport institutionnel.
            for point in data.get("quarterly_infc") or []:
                match = re.fullmatch(r"(\d{4})-T([1-4])", point.get("period") or "")
                if not match:
                    continue
                point_start, point_end = validate_period("TRIMESTRIEL", int(match[1]), None, int(match[2]))[:2]
                if point_start > as_of:
                    continue
                average = await DashboardRepository.infc_average_in_period(
                    db, start_date=point_start, end_date=point_end
                )
                point["value"] = str(average) if average is not None else None
        return {
            "type": kind, "title": TEMPLATES[kind][0], "period_label": period_label,
            "start": start.isoformat(), "end": end.isoformat(),
            "data_as_of": as_of.isoformat(), "provisional": end > today,
            "methodology": METHODOLOGY, "rows": readable_rows(data, kind, period_label, as_of=as_of),
        }

    @staticmethod
    async def generate(db: AsyncSession, *, kind: str, year: int, month: int | None, quarter: int | None, format: str, actor, request: Request):
        if format not in FORMATS:
            raise HTTPException(422, "Format de rapport invalide.")
        report = await BnecReportService.build(db, kind=kind, year=year, month=month, quarter=quarter)
        renderer = {"PDF": render_pdf, "XLSX": render_xlsx, "CSV": render_csv}[format]
        content = renderer(report["title"], report["period_label"], report["rows"])
        extension, _ = FORMATS[format]
        storage_name = f"{uuid4().hex}.{extension}"
        target = (storage_root() / storage_name).resolve()
        target.write_bytes(content)
        try:
            document = Document(
                type_document="RAPPORT_BNEC", nom_original=f"bilan-bnec-{kind.lower()}-{report['period_label'].replace('/', '-')}.{extension}",
                nom_stockage=storage_name, chemin_stockage=str(target), format=format,
                taille_octets=len(content), checksum=sha256(content).hexdigest(), version="1",
                ressource_type="RAPPORT_BNEC", confidentialite="INTERNE", source="GENERATION_SERVEUR",
                date_document=date.today(), depose_par_id=actor.user.id,
                date_depot=datetime.now(timezone.utc), statut_verification="NON_REQUIS", statut="ACTIF",
            )
            db.add(document)
            await db.flush()
            item = RapportGenere(
                code_modele=f"BNEC_{kind}", nom_modele=report["title"], categorie="BNEC",
                demandeur_id=actor.user.id, filtres={"type": kind, "year": year, "month": month, "quarter": quarter},
                sections={"source": "dashboard_serveur", "indicateurs": len(report["rows"])}, format=format,
                periode_debut=report["start"], periode_fin=report["end"], date_demande=date.today(),
                date_generation=date.today(), document_id=document.id, resultat=f"Bilan {report['period_label']} généré par le serveur.", statut="GENERE",
            )
            db.add(item)
            await db.flush()
            document.ressource_id = item.id
            await write_audit_event(
                db, action="REPORT_BNEC_GENERATE", categorie="REPORTING", resultat="SUCCES",
                utilisateur_id=actor.user.id, ressource_type="rapport_genere", ressource_id=item.id,
                adresse_ip=request.client.host if request.client else None,
                valeurs_apres={"modele": item.code_modele, "periode_debut": item.periode_debut, "periode_fin": item.periode_fin, "format": format, "document_id": str(document.id)},
            )
            await db.commit()
            await db.refresh(item)
            return GovernanceService.report_response(item)
        except Exception:
            await db.rollback()
            target.unlink(missing_ok=True)
            raise

    @staticmethod
    async def download(db: AsyncSession, report_id):
        item = await GovernanceService.require_report(db, report_id)
        if item.statut != "GENERE" or item.document_id is None:
            raise HTTPException(404, "Fichier de rapport indisponible.")
        document = await GovernanceRepository.get_active_document(db, item.document_id)
        if document is None or document.ressource_type != "RAPPORT_BNEC" or document.format not in FORMATS:
            raise HTTPException(404, "Fichier de rapport indisponible.")
        root = storage_root()
        path = Path(document.chemin_stockage or "").resolve()
        if root not in path.parents or not path.is_file():
            raise HTTPException(404, "Fichier de rapport introuvable.")
        return path, document.nom_original, FORMATS[document.format][1]
