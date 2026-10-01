from __future__ import annotations

import shutil
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.dml.color import RGBColor
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v3.docx"
REFERENCE = ROOT / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v2_mis_a_jour_final.docx"
OUTPUT = ROOT / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v3_mis_a_jour.docx"


def iter_paragraphs(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def all_paragraphs(doc):
    yield from iter_paragraphs(doc)
    seen = set()
    for section in doc.sections:
        for part in (section.header, section.footer):
            key = str(part.part.partname)
            if key in seen:
                continue
            seen.add(key)
            yield from iter_paragraphs(part)


def replace_text(paragraph, old: str, new: str) -> bool:
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    return False


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)
reference = Document(REFERENCE)

# Reprendre la palette institutionnelle de la version illustrée précédente.
for style in doc.styles:
    if style.name not in reference.styles or not hasattr(style, "font"):
        continue
    ref_style = reference.styles[style.name]
    if not hasattr(ref_style, "font"):
        continue
    ref_color = ref_style.font.color.rgb
    if ref_color is not None:
        style.font.color.rgb = RGBColor.from_string(str(ref_color))

# La V3 avait forcé chaque caractère en noir. Retirer ce format direct permet
# aux styles colorés de s'appliquer de nouveau, sans altérer les tailles, les
# graisses ou la disposition choisies par l'utilisateur.
for paragraph in all_paragraphs(doc):
    remove_paragraph_borders(paragraph)
    for run in paragraph.runs:
        if run.font.color.rgb is not None and str(run.font.color.rgb) == "000000":
            run.font.color.rgb = None

# Restaurer les couleurs directes des cartouches, légendes et en-têtes de
# tableaux qui ne reposaient pas uniquement sur un style Word.
reference_colors = defaultdict(list)
for paragraph in all_paragraphs(reference):
    text = paragraph.text.strip()
    if not text:
        continue
    colors = [str(run.font.color.rgb) for run in paragraph.runs if run.text.strip() and run.font.color.rgb]
    if colors and len(set(colors)) == 1:
        reference_colors[text].append(colors[0])

used = defaultdict(int)
for paragraph in all_paragraphs(doc):
    text = paragraph.text.strip()
    matches = reference_colors.get(text)
    if not matches:
        continue
    index = min(used[text], len(matches) - 1)
    used[text] += 1
    color = RGBColor.from_string(matches[index])
    for run in paragraph.runs:
        if run.text.strip():
            run.font.color.rgb = color

# Actualiser les repères de version visibles.
for paragraph in doc.paragraphs:
    replace_text(paragraph, "Version 3 - manuel opérationnel", "Version 3.1 - manuel opérationnel")

for table in doc.tables:
    if len(table.rows) >= 2 and table.cell(0, 0).text.strip() == "Version":
        table.cell(0, 1).text = "3.1 - guide opérationnel actualisé"
        if table.cell(1, 0).text.strip() == "Date":
            table.cell(1, 1).text = "30 septembre 2026"
        break

# Préciser la portée personnelle du centre des alertes, telle qu'elle est
# actuellement appliquée par l'interface et les permissions.
old_alert_step = "1. Ouvrir la liste des alertes."
new_alert_step = (
    "1. Ouvrir le Centre des alertes : l'onglet « Mes alertes » affiche par défaut "
    "les éléments destinés au compte connecté. Le « Registre général » n'est visible "
    "que si le rôle possède l'autorisation correspondante."
)
for paragraph in doc.paragraphs:
    if paragraph.text.strip() == old_alert_step:
        paragraph.text = new_alert_step
        paragraph.style = doc.styles["Normal"]
        break

# Consolider l'historique sans ajouter une ligne susceptible de créer une page
# blanche avant l'annexe RACI en orientation paysage.
for table in doc.tables:
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    if headers == ["Version", "Date", "Évolution"]:
        for row in table.rows[1:]:
            if row.cells[0].text.strip() == "3.0":
                row.cells[0].text = "3.1"
                row.cells[1].text = "30 septembre 2026"
                row.cells[2].text = (
                    "Mise à jour fonctionnelle complète, preuves par certificat, intégration BNEC, "
                    "règles publiées et palette institutionnelle rétablie."
                )
                break
        break

# Mettre à jour les pieds de page tout en conservant le champ automatique PAGE.
seen_parts = set()
for section in doc.sections:
    footer = section.footer
    part_key = str(footer.part.partname)
    if part_key in seen_parts:
        continue
    seen_parts.add(part_key)
    for paragraph in footer.paragraphs:
        replace_text(paragraph, "Version 1.0 - Août 2026", "Version 3.1 - Septembre 2026")
        replace_text(paragraph, "Version 3.0 - Septembre 2026", "Version 3.1 - Septembre 2026")
        for run in paragraph.runs:
            if run.text.strip():
                run.font.color.rgb = RGBColor(0x62, 0x7A, 0x72)

# Les nouveaux textes doivent suivre la même palette que leur contexte.
for paragraph in all_paragraphs(doc):
    if paragraph.style.name in {"Heading 1", "Heading 3"}:
        for run in paragraph.runs:
            if run.text.strip():
                run.font.color.rgb = RGBColor(0x06, 0x4A, 0x3A)
    elif paragraph.style.name == "Heading 2":
        for run in paragraph.runs:
            if run.text.strip():
                run.font.color.rgb = RGBColor(0x08, 0x76, 0x59)

doc.save(OUTPUT)
print(OUTPUT)
