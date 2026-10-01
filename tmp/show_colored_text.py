from pathlib import Path
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
doc = Document(ROOT / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v2_mis_a_jour_final.docx")


def walk(container, prefix="BODY"):
    for pidx, paragraph in enumerate(container.paragraphs):
        colors = sorted({str(r.font.color.rgb) for r in paragraph.runs if r.text.strip() and r.font.color.rgb})
        if colors:
            print(prefix, pidx, paragraph.style.name, colors, paragraph.text[:180])
    for tidx, table in enumerate(container.tables):
        for ridx, row in enumerate(table.rows):
            for cidx, cell in enumerate(row.cells):
                walk(cell, f"{prefix}/T{tidx}/R{ridx}/C{cidx}")


walk(doc)
for sidx, section in enumerate(doc.sections):
    walk(section.header, f"HEADER{sidx}")
    walk(section.footer, f"FOOTER{sidx}")
