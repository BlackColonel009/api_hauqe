from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.dml.color import RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v3_mis_a_jour.docx"

GREEN_DARK = "064A3A"
GREEN = "087659"
GREEN_TEXT = "183B32"
GREEN_PALE = "F2F7F5"
GRID = "B7DAD0"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), GRID)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def insert_paragraph_after(cursor, text: str = "", style: str = "Normal"):
    paragraph = doc.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    cursor.addnext(paragraph._p)
    paragraph.paragraph_format.keep_with_next = style.startswith("Heading")
    return paragraph._p, paragraph


def set_version_text(container, old: str, new: str) -> None:
    for paragraph in container.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


doc = Document(PATH)

if any(p.text.strip() == "Comprendre les niveaux de pilotage" for p in doc.paragraphs):
    raise RuntimeError("La rubrique des tableaux de bord existe déjà dans le guide.")

anchor_index, anchor = next(
    (index, paragraph)
    for index, paragraph in enumerate(doc.paragraphs)
    if paragraph.text.strip() == "Lire le tableau de bord"
)
intro = doc.paragraphs[anchor_index + 1]
intro.text = (
    "Le tableau de bord opérationnel sert au suivi quotidien. Les tableaux tactique, stratégique "
    "et annuel consolident ensuite les mêmes données à des horizons différents. Le Baromètre "
    "national observe la structure du registre, tandis que le tableau public ne présente que les "
    "indicateurs agrégés officiellement publiés."
)
intro.style = doc.styles["Normal"]

cursor = doc.paragraphs[anchor_index + 6]._p

cursor, _ = insert_paragraph_after(cursor, "Comprendre les niveaux de pilotage", "Heading 2")
cursor, _ = insert_paragraph_after(
    cursor,
    "Chaque tableau répond à une question différente. L'utilisateur commence par vérifier la période "
    "affichée et les filtres actifs, puis interprète l'indicateur avant d'ouvrir le dossier, le registre "
    "ou le rapport qui permet d'agir. Un indicateur attire l'attention ; il ne remplace ni le contrôle "
    "du dossier source ni la décision du rôle compétent.",
)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.style = doc.styles["Normal Table"]
widths = [Cm(2.6), Cm(3.5), Cm(4.6), Cm(6.0)]
headers = ["Espace", "Période et finalité", "Utilisateurs habilités", "Utilisation attendue"]
for index, cell in enumerate(table.rows[0].cells):
    cell.width = widths[index]
    cell.text = headers[index]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, GREEN_DARK)
    set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
repeat_header(table.rows[0])

dashboard_rows = [
    (
        "Opérationnel",
        "Quotidienne, sur 1 à 90 jours. Repérer les actions immédiates.",
        "Tous les comptes possédant DASHBOARDS.OPERATIONNEL.",
        "Filtrer, lire les alertes et échéances, ouvrir la ressource concernée et exporter le relevé si nécessaire.",
    ),
    (
        "Tactique",
        "Mensuelle. Suivre la chaîne de traitement et la qualité.",
        "Administrateur HAUQE, Direction technique, Point focal BNEC, Administrateur BNEC et Cellule de veille.",
        "Choisir l'année et le mois, comparer les étapes Collecte à Veille, puis ouvrir le rapport de période.",
    ),
    (
        "Stratégique",
        "Trimestrielle. Éclairer les décisions institutionnelles.",
        "Administrateur HAUQE et Direction technique.",
        "Analyser la couverture régionale, l'évolution de l'INFC, les constats, les risques majeurs et les recommandations.",
    ),
    (
        "Annuel",
        "Annuelle. Produire le bilan institutionnel.",
        "Administrateur HAUQE et Direction technique.",
        "Comparer les trimestres, la qualité, la gouvernance et la continuité, puis préparer le bilan annuel.",
    ),
    (
        "Baromètre national",
        "Période libre. Observer la structure nationale des certifications.",
        "Administrateur HAUQE, Direction technique, Point focal BNEC, Administrateur BNEC, Cellule de veille et Lecteur.",
        "Définir les dates et analyser les répartitions par région, secteur, norme, organisme, statut et classe SNCC.",
    ),
    (
        "Public",
        "Période officiellement publiée. Informer sans exposer de données confidentielles.",
        "Public, partenaires et tout visiteur.",
        "Consulter uniquement les indicateurs agrégés autorisés ; aucune action sur un dossier individuel n'est possible.",
    ),
    (
        "Veille",
        "Quotidienne après intégration BNEC. Suivre la vie des certifications.",
        "Cellule de veille et responsables autorisés.",
        "Traiter les dossiers de veille, alertes, relances, échéances, réponses et décisions à préparer.",
    ),
]

for row_index, values in enumerate(dashboard_rows, start=1):
    cells = table.add_row().cells
    keep_row_together(table.rows[-1])
    for index, (cell, value) in enumerate(zip(cells, values)):
        cell.width = widths[index]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if row_index % 2 == 0:
            set_cell_shading(cell, GREEN_PALE)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(GREEN_TEXT)
            if index == 0:
                run.font.bold = True

set_table_borders(table)
cursor.addnext(table._tbl)
cursor = table._tbl

cursor, _ = insert_paragraph_after(
    cursor,
    "Les droits effectivement attribués dans Administration restent la référence. Un menu absent ou un refus d'accès "
    "signifie que le compte ne possède pas la permission correspondante ; il ne faut pas utiliser le compte d'un autre agent.",
)

cursor, _ = insert_paragraph_after(cursor, "Utiliser le tableau de bord opérationnel", "Heading 3")
for text in (
    "1. Choisir une période comprise entre 1 et 90 jours.",
    "2. Appliquer les filtres de zone, secteur, norme ou organisme avant de comparer les valeurs.",
    "3. Lire les actions prioritaires, les certificats expirés ou proches de l'échéance et les certifications récentes.",
    "4. Cliquer sur une carte, une ligne ou le menu d'action pour ouvrir le registre, l'entreprise, la certification ou l'échéance concernée.",
    "5. Utiliser Exporter pour produire le relevé Excel du même état opérationnel.",
):
    cursor, _ = insert_paragraph_after(cursor, text)

cursor, _ = insert_paragraph_after(cursor, "Utiliser le tableau de bord tactique", "Heading 3")
for text in (
    "1. Sélectionner l'année et le mois à analyser.",
    "2. Examiner successivement Collecte, Vérification, FUCCS, Validation, Intégration et Veille afin de localiser les files en attente ou les retards.",
    "3. Comparer les volumes ouverts et clôturés, les délais et les données de qualité du mois.",
    "4. Ouvrir le rapport mensuel lorsque l'analyse doit être conservée ou transmise.",
):
    cursor, _ = insert_paragraph_after(cursor, text)

cursor, _ = insert_paragraph_after(cursor, "Utiliser le tableau de bord stratégique", "Heading 3")
for text in (
    "1. Choisir l'année et le trimestre.",
    "2. Comparer la couverture régionale et l'évolution de l'INFC sur la période.",
    "3. Lire les constats, les risques majeurs et les recommandations prioritaires produits à partir des données consolidées.",
    "4. Ouvrir le rapport trimestriel pour préparer la décision ou la réunion de pilotage.",
):
    cursor, _ = insert_paragraph_after(cursor, text)

cursor, _ = insert_paragraph_after(cursor, "Utiliser le tableau de bord annuel", "Heading 3")
for text in (
    "1. Sélectionner l'année du bilan.",
    "2. Comparer les certifications et l'INFC par trimestre.",
    "3. Examiner les volets Qualité, Gouvernance et Continuité afin d'identifier les progrès, les insuffisances et les incidents.",
    "4. Ouvrir le rapport annuel pour préparer le bilan institutionnel et les orientations de l'année suivante.",
):
    cursor, _ = insert_paragraph_after(cursor, text)

cursor, _ = insert_paragraph_after(cursor, "Utiliser le Baromètre national", "Heading 3")
for text in (
    "1. Définir la date de début et la date de fin de la période observée.",
    "2. Lire les nombres d'entreprises, de certifications, de certifications actives et l'INFC national moyen.",
    "3. Examiner les répartitions par région, secteur, norme, organisme, statut et classe SNCC.",
    "4. Utiliser ces résultats pour décrire la situation nationale ; revenir au registre pour contrôler un dossier individuel.",
):
    cursor, _ = insert_paragraph_after(cursor, text)

cursor, _ = insert_paragraph_after(cursor, "Consulter le tableau de bord public", "Heading 3")
for text in (
    "1. Vérifier la période et la référence de la publication affichée.",
    "2. Lire les indicateurs agrégés et l'avertissement de diffusion associé.",
    "3. En cas d'absence de données, vérifier qu'une règle et une publication institutionnelle sont effectivement publiées ; aucune donnée interne ne doit être ajoutée directement depuis cet écran.",
):
    cursor, _ = insert_paragraph_after(cursor, text)

cursor, _ = insert_paragraph_after(cursor, "Utiliser le tableau de veille", "Heading 3")
for text in (
    "1. Ouvrir l'espace Veille après l'intégration BNEC d'une certification.",
    "2. Examiner les cartes de suivi et la file des dossiers nécessitant une relance, un renouvellement ou une décision.",
    "3. Ouvrir le dossier source, enregistrer l'action, le responsable, l'échéance et la preuve de communication.",
    "4. Clôturer uniquement après réception du résultat attendu ou avec une justification motivée.",
):
    cursor, _ = insert_paragraph_after(cursor, text)

cursor, _ = insert_paragraph_after(cursor, "Relier les tableaux de bord au parcours", "Heading 2")
cursor, _ = insert_paragraph_after(
    cursor,
    "Les tableaux de bord exploitent les informations produites par les étapes du parcours. Ils ne créent pas un second dossier et ne modifient pas directement les décisions déjà enregistrées.",
)
for text in (
    "La collecte, la vérification, le contrôle FUCCS, les validations et l'intégration alimentent le tableau opérationnel.",
    "La consolidation mensuelle de ces étapes alimente le tableau tactique.",
    "Les résultats trimestriels, l'INFC et les risques alimentent le tableau stratégique.",
    "Les résultats de l'année, la qualité, la gouvernance et la continuité alimentent le tableau annuel.",
    "Les certifications intégrées et classées alimentent le Baromètre national et le suivi de veille.",
    "Seules les données agrégées approuvées dans Publications alimentent le tableau public.",
):
    cursor, _ = insert_paragraph_after(cursor, text, "List Bullet")

cursor, _ = insert_paragraph_after(cursor, "Points d'attention pour le pilotage", "Heading 2")
for text in (
    "Toujours vérifier la période, les filtres et la date de calcul avant de commenter un indicateur.",
    "Une hausse n'est pas automatiquement positive et une baisse n'est pas automatiquement négative ; interpréter la définition de l'indicateur et le dossier source.",
    "Un chiffre du tableau de bord ne vaut ni validation, ni classement, ni décision administrative.",
    "Les données publiques doivent provenir du circuit Règles et codification, demande, approbation et publication.",
):
    cursor, _ = insert_paragraph_after(cursor, text, "List Bullet")

# Mise à jour de la version visible sans créer une nouvelle ligne de tableau.
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.text = run.text.replace("Version 3.1 - manuel opérationnel", "Version 3.2 - manuel opérationnel")

doc.tables[0].cell(0, 1).text = "3.2 - guide opérationnel actualisé"
doc.tables[0].cell(1, 1).text = "1er octobre 2026"

for table_candidate in doc.tables:
    if [cell.text.strip() for cell in table_candidate.rows[0].cells] != ["Version", "Date", "Évolution"]:
        continue
    for row in table_candidate.rows[1:]:
        if row.cells[0].text.strip() == "3.1":
            row.cells[0].text = "3.2"
            row.cells[1].text = "1er octobre 2026"
            row.cells[2].text = (
                "Mise à jour fonctionnelle, palette institutionnelle et guide des tableaux de bord "
                "opérationnel, tactique, stratégique, annuel, national, public et de veille."
            )
            break
    break

seen_footer_parts = set()
for section in doc.sections:
    footer = section.footer
    key = str(footer.part.partname)
    if key in seen_footer_parts:
        continue
    seen_footer_parts.add(key)
    set_version_text(footer, "Version 3.1 - Septembre 2026", "Version 3.2 - Octobre 2026")
    for paragraph in footer.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                run.font.color.rgb = RGBColor(0x62, 0x7A, 0x72)

doc.save(PATH)
print(PATH)
