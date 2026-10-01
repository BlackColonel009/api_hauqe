from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
path = root / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v3_mis_a_jour.docx"

dash_items = {
    "La collecte, la vérification, le contrôle FUCCS, les validations et l'intégration alimentent le tableau opérationnel.",
    "La consolidation mensuelle de ces étapes alimente le tableau tactique.",
    "Les résultats trimestriels, l'INFC et les risques alimentent le tableau stratégique.",
    "Les résultats de l'année, la qualité, la gouvernance et la continuité alimentent le tableau annuel.",
    "Les certifications intégrées et classées alimentent le Baromètre national et le suivi de veille.",
    "Seules les données agrégées approuvées dans Publications alimentent le tableau public.",
    "Toujours vérifier la période, les filtres et la date de calcul avant de commenter un indicateur.",
    "Une hausse n'est pas automatiquement positive et une baisse n'est pas automatiquement négative ; interpréter la définition de l'indicateur et le dossier source.",
    "Un chiffre du tableau de bord ne vaut ni validation, ni classement, ni décision administrative.",
    "Les données publiques doivent provenir du circuit Règles et codification, demande, approbation et publication.",
}

doc = Document(path)
matched = 0
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text in dash_items:
        paragraph.style = doc.styles["Normal"]
        paragraph.text = "- " + text
        matched += 1
    elif text.startswith("- ") and text[2:] in dash_items:
        paragraph.style = doc.styles["Normal"]
        matched += 1

# Maintenir le tableau dans la largeur utile de la page portrait (16,6 cm).
dashboard_table = next(
    table for table in doc.tables
    if table.cell(0, 0).text.strip() == "Espace"
    and table.cell(0, 1).text.strip() == "Période et finalité"
)
widths = [Cm(2.5), Cm(3.4), Cm(4.5), Cm(6.1)]
for row in dashboard_table.rows:
    for cell, width in zip(row.cells, widths):
        cell.width = width
for grid_col, width in zip(dashboard_table._tbl.tblGrid.gridCol_lst, widths):
    grid_col.set(qn("w:w"), str(int(width.twips)))

# Demander à Word de recalculer les champs (sommaire, listes et pagination) à l'ouverture.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.save(path)
assert matched == len(dash_items), (matched, len(dash_items))
with ZipFile(path) as package:
    assert package.testzip() is None
print(path)
print("dash_items", matched)
