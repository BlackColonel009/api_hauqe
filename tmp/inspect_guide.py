from pathlib import Path
from collections import Counter

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
FILES = {
    "v2": ROOT / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v2_mis_a_jour_final.docx",
    "v3": ROOT / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v3.docx",
}


def rgb_of(run):
    value = run.font.color.rgb
    return str(value) if value is not None else "AUTO"


for label, path in FILES.items():
    doc = Document(path)
    print(f"\n=== {label}: {path.name} ===")
    for style_name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Normal"]:
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        color = style.font.color.rgb
        print(f"STYLE {style_name}: color={color}, size={style.font.size}, bold={style.font.bold}")

    for style_name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        colors = Counter()
        samples = []
        for idx, paragraph in enumerate(doc.paragraphs):
            if paragraph.style.name != style_name or not paragraph.text.strip():
                continue
            for run in paragraph.runs:
                if run.text.strip():
                    colors[rgb_of(run)] += 1
            if len(samples) < 5:
                samples.append((idx, paragraph.text[:90], [(r.text[:30], rgb_of(r)) for r in paragraph.runs if r.text.strip()]))
        print(f"RUN COLORS {style_name}: {colors}")
        for sample in samples:
            print("  ", sample)

    keywords = [
        "échéance", "Intégration BNEC", "plan d'intégration", "plan d’intégration",
        "codification", "séquence", "complément", "Historique des versions",
        "courriels du système", "Mes alertes", "menu mobile", "formulaire modal",
        "preuve", "statut ?", "assistant", "précréation",
    ]
    print("KEYWORD PARAGRAPHS")
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text and any(keyword.casefold() in text.casefold() for keyword in keywords):
            print(f"P{idx:03d} [{paragraph.style.name}] {text[:260]}")

    print("TABLES MATCHING KEYWORDS")
    for tidx, table in enumerate(doc.tables):
        text = " | ".join(cell.text.replace("\n", " / ").strip() for row in table.rows for cell in row.cells)
        if any(keyword.casefold() in text.casefold() for keyword in keywords):
            print(f"T{tidx:03d} {text[:700]}")
