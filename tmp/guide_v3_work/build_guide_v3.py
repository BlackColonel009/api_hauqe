from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\hp\Documents\APK WEB Projets R.1.3.5 et R1.4.3")
SOURCE = ROOT / "output/docx/Guide_global_utilisation_SNGSC_HAUQE_v2_mis_a_jour_final.docx"
DEST = ROOT / "output/docx/Guide_global_utilisation_SNGSC_HAUQE_v3.docx"

doc = Document(SOURCE)
paragraphs = list(doc.paragraphs)

replacements = {
    9: "Guide global d'utilisation du SNGSC",
    10: "Version 3 - manuel opérationnel pour les agents habilités",
    43: "Le parcours d'un dossier suit six étapes : collecte, vérification documentaire, contrôle FUCCS, validation N1, validation N2 et intégration BNEC. Le bloc « Suivi du dossier » indique l'étape atteinte, la prochaine action et les blocages. Après intégration, la certification peut être suivie dans la veille, le scoring, l'INFC et le SNCC. Une étape achevée ne donne pas automatiquement le droit d'exécuter la suivante : les permissions et l'affectation restent applicables.",
    49: "Ouvrir le dossier et lire le bloc « Suivi du dossier », son statut et l'étape suivante.",
    63: "En cas de correction après soumission ou intégration, utiliser la révision ou la réouverture autorisée ; ne jamais créer un second dossier pour contourner le parcours.",
    82: "Ouvrir le lien reçu par courriel dans les 3 minutes suivant son émission ; passé ce délai, demander un nouveau lien.",
    83: "Définir un nouveau mot de passe, puis se reconnecter : les sessions précédentes sont révoquées et une notification de sécurité est envoyée. Le nouveau mot de passe n'est jamais communiqué par courriel.",
    99: "Chaque utilisateur peut tenir à jour son identité et ses préférences. Les administrateurs habilités disposent aussi d'un réglage de réception des courriels fonctionnels, compte par compte. Cette préférence n'enlève aucune notification personnelle dans l'application et ne modifie pas les droits métier.",
    120: "Marquer la notification comme lue après avoir ouvert le dossier ou l'action concernée.",
    156: "Les alertes et les échéances dépendent des dates et paramètres enregistrés, et non d'un texte fixe. Un certificat déjà expiré doit apparaître dans la catégorie des expirés, même s'il n'est pas dans la fenêtre des prochains jours.",
    159: "Dans « Configurer les rappels », définir l'horizon et les destinataires d'une échéance, puis vérifier le plan enregistré dans le détail. Les rappels de l'agent et les avis aux administrateurs peuvent être désactivés selon les droits disponibles.",
    170: "Rechercher d'abord par raison sociale, identifiant HAUQE ou localisation. Les identifiants juridiques facultatifs ne sont pas nécessaires pour retrouver une entreprise.",
    173: "Enregistrer, puis contrôler les contacts, les offres, les destinations, les sites et les documents dans la fiche entreprise.",
    183: "Dans la collecte, l'agent peut rechercher l'organisme certificateur et, si aucun résultat fiable n'existe, le précréer dans un formulaire dédié. Il renseigne l'identité, les contacts et les données d'accréditation connues ; le statut reste « À vérifier » jusqu'au rapprochement métier.",
    184: "Commencer par taper le nom de l'organisme ; sélectionner une fiche existante si elle correspond.",
    185: "Si aucun résultat fiable n'apparaît, utiliser « Précréer », puis saisir le nom officiel et les renseignements disponibles, sans inventer un accréditeur ou une date.",
    186: "Valider la précréation et vérifier que l'organisme est sélectionné sur la certification déclarée.",
    197: "Une campagne regroupe plusieurs missions ; chaque mission peut regrouper plusieurs fiches de collecte d'entreprises. L'administrateur affecte un ou plusieurs agents à la mission. Chaque fiche est attribuée à l'agent qui la crée ; les autres agents affectés peuvent ouvrir leur propre collecte sur la même mission, sans modifier son brouillon. La déclaration terrain ne devient une certification officielle qu'après les contrôles, les deux validations et l'intégration BNEC.",
    204: "Cliquer sur « Nouvelle campagne » ; le code proposé est incrémenté pour l'année courante. Renseigner le nom, l'objet, l'objectif et la période prévisionnelle.",
    211: "La mission, rattachée à une campagne, précise la référence, la zone, les dates et l'objet communs à plusieurs collectes. Sa référence est proposée automatiquement à la création. L'administrateur ou le rôle habilité peut y affecter plusieurs agents, puis compléter cette affectation sans réviser les fiches existantes.",
    212: "Depuis « Collecte et contrôle », ouvrir la campagne, puis créer une mission distincte ou sélectionner une mission déjà créée.",
    213: "Vérifier la référence proposée, la zone administrative, les dates, la priorité et l'objet ; précréer la zone si elle manque et si votre rôle le permet.",
    214: "Sélectionner tous les agents autorisés à collecter dans cette mission, puis enregistrer l'affectation.",
    215: "Dans la mission, utiliser « Nouvelle collecte » pour ouvrir une fiche propre à l'entreprise visitée. Une autre entreprise exige une autre fiche, pas une autre mission.",
    219: "Rechercher l'entreprise existante avant toute précréation ; ne pas créer une deuxième fiche pour une simple variante de nom.",
    220: "Si elle n'existe pas, utiliser le formulaire de précréation pour l'identité, la zone, l'adresse, la localisation et les coordonnées disponibles. Le code temporaire est proposé selon l'année courante ; le code BNEC définitif vient à l'intégration.",
    226: "Ajouter chaque produit ou service avec sa catégorie, sa description et ses marchés visés. Les informations exploitables peuvent alimenter la fiche entreprise à la soumission ou compléter ses champs vides lors d'une intégration ultérieure.",
    227: "Ajouter chaque certification avec son numéro, sa norme, sa portée, ses dates et l'organisme certificateur sélectionné ou précréé. Une certification différente doit être une ligne distincte.",
    232: "Les preuves générales de la fiche et les preuves propres à chaque certificat ont des rattachements différents. Avant la soumission, vérifier que chaque fichier sélectionné apparaît dans la bonne rubrique, que le déclarant et le consentement sont renseignés, et que les dates et numéros concordent.",
    233: "Dans chaque certification déclarée, ajouter sa preuve spécifique et contrôler le nom du fichier affiché ; la disponibilité de copie passe alors à « Oui ».",
    234: "Dans « Preuves et observations », ajouter seulement les justificatifs généraux de la collecte. Une preuve générale n'est pas attribuée automatiquement à tous les certificats.",
    239: "Une mission peut avoir plusieurs collectes, mais chaque brouillon reste propre à son agent responsable. Ne pas réutiliser la fiche d'une autre entreprise.",
    244: "Dans un brouillon, « Enregistrer » conserve la saisie ; « Réinitialiser » efface les champs de la fiche courante après confirmation, sans modifier la campagne, la mission ni les autres collectes. Après soumission, passer par une révision tracée.",
    254: "Le vérificateur consulte la collecte et les pièces liées, y compris les preuves propres aux certifications déclarées. Il confronte les déclarations aux documents et aux références et consigne les anomalies, les confirmations externes et sa synthèse. L'affectation des vérificateurs demeure une action distincte réservée aux rôles habilités.",
    256: "Comparer chaque information de la fiche aux justificatifs généraux et aux preuves du certificat concerné ; ouvrir ou télécharger la pièce si nécessaire.",
    268: "La réouverture d'une vérification dépend de l'avancement réel du dossier. Le modal explique l'impact avant confirmation : sans contrôle FUCCS, la reprise est possible avec motif ; après contrôle ou validation, les étapes aval doivent être reprises ou la réouverture est interdite ; après intégration BNEC, créer une révision de collecte plutôt que modifier directement le dossier intégré.",
    269: "Cliquer sur « Réouvrir », lire l'impact affiché et renseigner un motif si l'action reste autorisée.",
    271: "Si la réouverture est interdite après intégration, utiliser le lien vers la révision de la collecte liée.",
    296: "L'intégration est ouverte seulement après la validation N2 favorable. Son aperçu présente les rapprochements, les preuves spécifiques, les éventuels compléments de champs vides et les codes proposés. Il faut lire les blocages avant de confirmer : l'aperçu n'est pas la réservation définitive des codes.",
    299: "Contrôler le titulaire, l'organisme, les numéros et dates, les preuves propres à chaque certification, ainsi que les codes proposés. Dans un même périmètre de codification, deux certificats du plan doivent recevoir des séquences distinctes ; des périmètres différents peuvent réutiliser un même suffixe.",
    312: "L'intégration produit les certifications officielles et leurs liens documentaires, sans distribuer une preuve générale de collecte à tous les certificats. Elle peut compléter seulement certains champs vides d'anciens dossiers ; elle ne remplace pas une valeur déjà renseignée.",
    313: "Les décisions N1 et N2 sont prononcées par deux utilisateurs distincts, chacun habilité pour son niveau. Une affectation ou une notification ne vaut pas décision.",
    320: "Le score FUCCS, la classification de l'entreprise, l'INFC de chaque certification et le classement SNCC restent quatre résultats distincts. Les trois derniers peuvent être proposés automatiquement à partir des données du parcours et des règles publiées, avec un rapport des manquements. L'utilisateur habilité doit examiner et valider le résultat avant de le considérer comme acquis.",
    326: "Lancer « Évaluer automatiquement », puis lire le détail des données prises en compte, du niveau proposé et des manquements.",
    327: "Corriger le dossier source si nécessaire ; valider l'évaluation seulement lorsque le résultat est justifié.",
    331: "L'INFC mesure la fiabilité d'une certification sur 100. Le moteur utilise les critères et leurs poids dans le modèle publié ; le poids fixe l'importance relative du critère dans la note totale. Les notes proposées et les manquements doivent être contrôlés avant validation.",
    334: "Lancer « Calculer automatiquement » ou « Recalculer automatiquement », puis ouvrir le rapport expliquant les notes et les manquements.",
    338: "Le classement SNCC présente trois dimensions indépendantes : classe A+ à D, statut administratif VA, RE, SU, RT, EX ou VE, et risque R1 à R5. Il se fonde sur la matrice publiée et sur la situation réelle du certificat ; l'expiration, le retrait ou la suspension ne doivent pas être masqués par une note favorable.",
    340: "Le moteur peut proposer automatiquement une classe, un statut et un risque à partir de l'INFC validé et de la matrice SNCC publiée. Il applique d'abord les situations prioritaires : expiré (EX), retiré (RT), suspendu (SU), puis non authentifié mais encore valide (VE). Sinon, la matrice de score détermine la proposition VA ou RE. L'utilisateur contrôle les sources et confirme le résultat motivé.",
    343: "Vérifier qu'une matrice SNCC publiée et un INFC validé existent ; consulter la certification, ses dates, son authenticité et les décisions antérieures.",
    344: "Cliquer sur « Calculer automatiquement » pour obtenir la proposition et le détail des manquements ; utiliser les actions de classement ou de reclassement autorisées pour confirmer.",
    345: "Contrôler séparément la classe, le statut administratif et le risque proposés. En cas d'écart avec le dossier, corriger la source plutôt que forcer un code.",
    354: "La classification entreprise, l'INFC et le SNCC demeurent dans trois espaces séparés, mais les règles publiées et les données du parcours alimentent leurs propositions automatiques.",
    372: "Saisir un destinataire explicite, une adresse valide, le contenu et la date d'envoi. Sélectionner les informations métier utiles à partager : entreprise, mission, certification, numéro, norme, portée, dates et produits, selon le cas.",
    375: "Pour un envoi prévu aujourd'hui, vérifier son état dans la file de notifications. « Envoyé au relais SMTP » confirme l'acceptation technique, pas la lecture par le destinataire.",
    400: "Suivre les étapes du formulaire : identité professionnelle, mot de passe initial, puis rôles initiaux. Vérifier les permissions réelles avant de créer le compte.",
    407: "Créer un brouillon ; utiliser le préremplissage proposé si nécessaire, puis lire son contenu et le résultat de la vérification avant publication.",
    429: "Les rapports, le journal d'audit et les revues de qualité servent au contrôle et à la traçabilité. Le module de sauvegarde suit les politiques et les opérations ; une sauvegarde complète ou une restauration réelle doit être vérifiée avec l'administration technique du serveur, notamment pour PostgreSQL et les documents.",
    447: "Le module présente les politiques et l'historique des opérations. Avant de compter sur une sauvegarde pour reprendre une production, vérifier qu'elle couvre la base, les comptes, l'historique et les fichiers, puis tester la restauration dans un environnement isolé avec l'administrateur serveur.",
    448: "Ouvrir « Sauvegardes » et examiner le périmètre et le statut de la politique active.",
    449: "Lancer ou planifier l'opération uniquement avec les droits requis et selon la procédure d'exploitation du serveur.",
    450: "Contrôler le résultat, l'intégrité et la disponibilité des fichiers sauvegardés.",
    451: "Documenter et tester périodiquement une restauration sur un environnement isolé avant toute intervention en production.",
}

for index, text in replacements.items():
    assert index < len(paragraphs) and paragraphs[index].text.strip(), index
    paragraphs[index].text = text


def insert_after(anchor, style, text):
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    from docx.text.paragraph import Paragraph

    p = Paragraph(element, anchor._parent)
    p.style = style
    p.add_run(text)
    return p


additions = {
    94: [
        ("List Bullet", "Le courriel de réinitialisation indique explicitement la durée de validité de 3 minutes. Si le lien expire, recommencer la demande ; ne pas réutiliser l'ancien lien."),
    ],
    126: [
        ("Heading 2", "Courriels fonctionnels des agents"),
        ("Normal", "Dans Profil puis Notifications, l'administrateur HAUQE ou BNEC habilité ouvre la liste des utilisateurs. Tous les comptes sont autorisés à recevoir les courriels fonctionnels par défaut. Il peut retirer ou rétablir un compte de cette liste et enregistrer. Les notifications dans l'application continuent d'arriver ; les courriels de sécurité du compte ne sont pas désactivés par ce réglage."),
    ],
    160: [
        ("List Bullet", "Dans le calendrier, les boutons précédent, aujourd'hui et suivant changent immédiatement de période. Cliquer sur une date ouvre les échéances de ce jour ; le registre consolidé permet de retrouver aussi les échéances hors période affichée."),
    ],
    192: [
        ("List Bullet", "Dans la fiche d'un organisme, la rubrique Documents montre ses pièces propres et, en lecture agrégée, les preuves actives des certifications qu'il a délivrées. Les fichiers ne sont pas copiés une seconde fois."),
        ("List Bullet", "Une certification « À vérifier » exige une preuve active et la confirmation de son authenticité ; déposer une copie ne suffit pas à confirmer celle-ci. Une date expirée reste signalée comme telle."),
    ],
    237: [
        ("Heading 2", "Reprendre ou réviser une collecte"),
        ("Normal", "Un brouillon peut être repris par son agent responsable. Les boutons de retrait d'offre ou de certification agissent sur la ligne choisie seulement. Une fiche déjà soumise ne se vide pas : une révision crée un nouveau brouillon, conserve l'historique et reprend les références documentaires actives à vérifier de nouveau."),
    ],
    315: [
        ("List Bullet", "Le code affiché dans le plan est une proposition. Si une autre intégration réserve entre-temps la même séquence, le code définitif est recalculé à la confirmation."),
        ("List Bullet", "Une certification avec un même numéro mais une autre norme ou un autre organisme peut bloquer le rapprochement automatique. Ne pas corriger ce conflit en inventant un numéro."),
    ],
    355: [
        ("List Bullet", "Sur la fiche certification, lire « Pourquoi ce statut ? » pour connaître les dates dépassées, les pièces non confirmées et l'action attendue. Une classe favorable ne masque jamais un statut EX, RT ou SU."),
    ],
    388: [
        ("List Bullet", "Après chaque étape majeure du parcours, les acteurs concernés reçoivent une notification personnelle et, si leur réception est active, un courriel contextualisé. Les rôles cumulés ne génèrent pas de doublons pour le même utilisateur."),
    ],
    424: [
        ("List Bullet", "Dans Règles et codification, les matrices publiées pilotent les propositions futures de classification, d'INFC et de SNCC. Préremplir un brouillon n'est pas publier une règle : vérifier le résultat, puis utiliser l'action de publication distincte."),
    ],
}

for index, entries in additions.items():
    anchor = paragraphs[index]
    for style, text in entries:
        anchor = insert_after(anchor, style, text)

# The source had one numbering sequence running from the first chapter to the
# last. Each procedure now starts at 1 under its own subheading.
step_number = 0
for p in doc.paragraphs:
    if p.style.name in ("Heading 1", "Heading 2", "Heading 3"):
        step_number = 0
    elif p.style.name == "List Number":
        step_number += 1
        p.style = "Normal"
        p.text = f"{step_number}. {p.text}"
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.19)
        p.paragraph_format.space_after = Inches(0.03)

# Keep the V2 visuals and tables, while making every authored text run monochrome.
for style in doc.styles:
    if style.type == 1:
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.highlight_color = None
for p in doc.paragraphs:
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.highlight_color = None
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.font.highlight_color = None

doc.tables[0].cell(0, 1).text = "3.0 - guide opérationnel actualisé"
doc.tables[0].cell(1, 1).text = "Septembre 2026"
sample = doc.tables[78]
sample.cell(5, 2).text = "Non renseignés dans cet exemple ; ces identifiants restent facultatifs et n'empêchent pas la collecte."
sample.cell(12, 2).text = "Preuve du certificat : certificat_iso22000_exemple.pdf ; preuves générales : photo_site_exemple.jpg et registre_production_exemple.pdf."
sample._tbl.remove(sample.rows[-1]._tr)
history = doc.tables[79]
history.cell(1, 0).text = "2.0"
history.cell(1, 1).text = "Septembre 2026"
history.cell(1, 2).text = "Guide illustré et détaillé du parcours SNGSC."
history._tbl.append(deepcopy(history.rows[1]._tr))
row = history.rows[-1]
row.cells[0].text = "3.0"
row.cells[1].text = "29 septembre 2026"
row.cells[2].text = "Collectes multiples par mission, preuves par certificat, parcours et notifications, intégration BNEC, calculs automatiques et règles publiées."

for table in (doc.tables[0], history):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)

for section in doc.sections:
    for p in section.header.paragraphs:
        for run in p.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)

# The V2 footer had a literal "1", not a working page field. Replace it with
# a real Word PAGE field in the shared footer.
footer = doc.sections[0].footer.paragraphs[0]
footer.text = "Diffusion interne - Version 3.0 - Septembre 2026  |  Page "
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(0, 0, 0)
field = footer.add_run()
for tag, value in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
    if tag is not None:
        item = OxmlElement("w:fldChar")
        item.set(qn("w:fldCharType"), tag)
    else:
        item = OxmlElement("w:instrText" if value.strip() == "PAGE" else "w:t")
        item.text = value
        if value.strip() == "PAGE":
            item.set(qn("xml:space"), "preserve")
    field._r.append(item)
field.font.color.rgb = RGBColor(0, 0, 0)

# Preserve the original table and paragraph geometry while replacing green or
# orange fills with neutral grays. Screenshots remain unchanged as evidence of
# the actual interface.
parts = [doc.element, doc.styles.element]
for section in doc.sections:
    parts.extend((section.header._element, section.footer._element))
for part in parts:
    for element in part.iter():
        if element.tag == qn("w:shd"):
            current = element.get(qn("w:fill"), "").upper()
            if current not in ("", "AUTO", "FFFFFF"):
                element.set(qn("w:fill"), "E4E4E4" if current in ("087659", "064A3A") else "F2F2F2")
        elif element.tag == qn("w:color"):
            if element.get(qn("w:val"), "").lower() not in ("", "auto"):
                element.set(qn("w:val"), "000000")

doc.core_properties.title = "Guide global d'utilisation du SNGSC - version 3"
doc.core_properties.subject = "Manuel opérationnel HAUQE Certif"
doc.save(DEST)
print(DEST)
