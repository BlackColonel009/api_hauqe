from pathlib import Path

from docx import Document
from docx.dml.color import RGBColor


root = Path(__file__).resolve().parents[1]
path = root / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v3_mis_a_jour.docx"
doc = Document(path)
seen = set()
for section in doc.sections:
    footer = section.footer
    key = str(footer.part.partname)
    if key in seen:
        continue
    seen.add(key)
    for paragraph in footer.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace(
                "Version 3.0 - Septembre 2026",
                "Version 3.1 - Septembre 2026",
            )
            if run.text.strip():
                run.font.color.rgb = RGBColor(0x62, 0x7A, 0x72)

for table in doc.tables:
    if [cell.text.strip() for cell in table.rows[0].cells] != ["Version", "Date", "Évolution"]:
        continue
    rows_31 = [row for row in table.rows[1:] if row.cells[0].text.strip() == "3.1"]
    for row in rows_31:
        table._tbl.remove(row._tr)
    target = None
    for row in table.rows[1:]:
        if row.cells[0].text.strip() == "3.0":
            target = row
    if target is not None:
        target.cells[0].text = "3.1"
        target.cells[1].text = "30 septembre 2026"
        target.cells[2].text = (
            "Mise à jour fonctionnelle complète, preuves par certificat, intégration BNEC, "
            "règles publiées et palette institutionnelle rétablie."
        )
    break
doc.save(path)
