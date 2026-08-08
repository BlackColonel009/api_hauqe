from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "docx" / "Guide_global_utilisation_SNGSC_HAUQE_v2.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

GREEN = "087659"
DARK = "064A3A"
MINT = "EAF6F1"
PALE = "F7FBF9"
LINE = "C9DED6"
ORANGE = "FFF2D9"
RED = "FBEAEA"
INK = "183B32"
MUTED = "627A72"


def set_font(run, name="Aptos", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color=LINE):
    tc_pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), color)
        b.append(e)
    tc_pr.append(b)


def cell_margin(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcPr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=INK, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margin(cell); borders(cell)


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def add_field(paragraph, instruction):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    paragraph._p.append(fld)


def add_page_break(doc):
    doc.add_page_break()


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, emphasis=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(7)
    if emphasis and emphasis in text:
        before, after = text.split(emphasis, 1)
        p.add_run(before)
        r = p.add_run(emphasis); set_font(r, color=GREEN, bold=True)
        p.add_run(after)
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_steps(doc, steps):
    for step in steps:
        p = doc.add_paragraph(style="List Number")
        p.add_run(step)


def add_callout(doc, title, text, tone="info"):
    fill = MINT if tone == "info" else ORANGE if tone == "attention" else RED
    color = DARK if tone != "warning" else "A33A3A"
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(t, [16.5])
    cell = t.cell(0, 0); shade(cell, fill); cell_margin(cell, 150, 170, 150, 170)
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title.upper()); set_font(r, size=8.5, color=color, bold=True)
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text); set_font(r2, size=9, color=INK)
    borders(cell, "A9CFC0" if tone == "info" else "E8C47D")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_kv_table(doc, rows):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(t, [4.1, 12.4])
    for i, (label, value) in enumerate(rows):
        cells = t.add_row().cells
        shade(cells[0], MINT); shade(cells[1], "FFFFFF" if i % 2 == 0 else PALE)
        set_cell_text(cells[0], label, True, DARK)
        set_cell_text(cells[1], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_matrix(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(t, widths)
    for cell, header in zip(t.rows[0].cells, headers):
        shade(cell, GREEN); set_cell_text(cell, header, True, "FFFFFF", 8.7)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            shade(cell, "FFFFFF" if i % 2 == 0 else PALE)
            set_cell_text(cell, value, False, INK, 8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure_placeholder(doc, number, title, note):
    add_heading(doc, f"Figure {number} - {title}", 3)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(t, [16.5])
    cell = t.cell(0, 0); shade(cell, "F2F7F5"); cell_margin(cell, 850, 180, 850, 180)
    cell.text = ""
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CAPTURE D'ÉCRAN À INSÉRER"); set_font(r, size=11, color=GREEN, bold=True)
    p2 = cell.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Masquer les données nominatives, identifiants, adresses électroniques et informations sensibles.")
    set_font(r2, size=8.5, color=MUTED, italic=True)
    borders(cell, "9EC8B9")
    p3 = doc.add_paragraph(); p3.paragraph_format.space_before = Pt(3); p3.paragraph_format.space_after = Pt(10)
    r3 = p3.add_run("Légende à compléter : "); set_font(r3, size=8.8, color=DARK, bold=True)
    r4 = p3.add_run(note); set_font(r4, size=8.8, color=MUTED)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2); normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7); normal.paragraph_format.line_spacing = 1.18
    for level, size, color in ((1, 17, DARK), (2, 13, GREEN), (3, 10.5, DARK)):
        st = styles[f"Heading {level}"]
        st.font.name = "Aptos Display"; st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True; st.paragraph_format.space_before = Pt(16 if level == 1 else 11)
        st.paragraph_format.space_after = Pt(6); st.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        st = styles[style_name]; st.font.name = "Aptos"; st.font.size = Pt(10)
        st.paragraph_format.space_after = Pt(3); st.paragraph_format.line_spacing = 1.12


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("HAUQE  |  SNGSC - Guide global d'utilisation")
    set_font(r, size=8.2, color=GREEN, bold=True)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Diffusion interne - Version 1.0 - Août 2026  |  Page ")
    set_font(r, size=8, color=MUTED)
    add_field(p, "PAGE")


def add_chapter(doc, number, title, purpose, roles, prerequisites, procedures, points, figure_title, figure_note):
    add_page_break(doc)
    add_heading(doc, f"{number}. {title}", 1)
    add_callout(doc, "But du chapitre", purpose)
    add_kv_table(doc, [("Utilisateurs concernés", roles), ("Prérequis", prerequisites)])
    for proc_title, steps, result in procedures:
        add_heading(doc, proc_title, 2)
        add_steps(doc, steps)
        add_callout(doc, "Résultat attendu", result)
    add_heading(doc, "Points d'attention", 2)
    add_bullets(doc, points)
    add_figure_placeholder(doc, number, figure_title, figure_note)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21); section.page_height = Cm(29.7)
section.top_margin = Cm(2.1); section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
section.header_distance = Cm(0.8); section.footer_distance = Cm(0.8)
setup_styles(doc); add_header_footer(section)

# Couverture
for _ in range(6): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("HAUQE"); set_font(r, size=16, color=GREEN, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SYSTÈME NATIONAL DE GESTION ET DE SUIVI DES CERTIFICATIONS"); set_font(r, size=11, color=DARK, bold=True)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Guide global d'utilisation"); set_font(r, name="Aptos Display", size=29, color=DARK, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Manuel opérationnel pour les agents habilités"); set_font(r, size=14, color=GREEN)
for _ in range(7): doc.add_paragraph()
add_kv_table(doc, [("Version", "1.0 - version de travail structurée"), ("Date", "Août 2026"), ("Diffusion", "Interne HAUQE - utilisateurs habilités"), ("Statut", "Document modifiable - captures à insérer par la HAUQE")])
add_page_break(doc)

# Front matter
add_heading(doc, "À propos de ce guide", 1)
add_body(doc, "Ce manuel accompagne l'utilisation fonctionnelle du SNGSC. Il décrit les actions attendues, les contrôles à réaliser et les résultats à vérifier. Les écrans peuvent varier selon le rôle attribué à l'utilisateur.")
add_callout(doc, "Règle de sécurité", "Ne partagez jamais votre mot de passe, votre code MFA, un code de récupération, un lien de réinitialisation ou une capture contenant des données sensibles.", "attention")
add_heading(doc, "Fiche documentaire", 2)
add_kv_table(doc, [("Titre", "Guide global d'utilisation du SNGSC"), ("Propriétaire", "HAUQE"), ("Utilisateurs", "Agents autorisés, superviseurs, administrateurs et responsables métier"), ("Mise à jour", "À réviser après chaque évolution fonctionnelle majeure")])
add_heading(doc, "Comment utiliser ce manuel", 2)
add_bullets(doc, ["Lire d'abord les chapitres 1 à 3 pour maîtriser l'accès, la sécurité et la navigation.", "Consulter ensuite le chapitre correspondant au rôle et au traitement confié.", "Insérer les captures dans les cadres prévus, puis compléter leurs légendes avec les repères visibles.", "En cas de doute, privilégier la consultation de l'historique avant toute modification."])
add_heading(doc, "Rôles de référence", 2)
add_matrix(doc, ["Rôle", "Responsabilité principale", "Modules usuels"], [
    ("Agent de collecte", "Constitue les fiches et rassemble les preuves.", "Collecte, entreprises, documents"),
    ("Vérificateur", "Contrôle la cohérence documentaire et les anomalies.", "Vérification, documents, entreprises"),
    ("Contrôleur FUCCS", "Applique la grille de contrôle et consigne les constats.", "FUCCS, documents, certifications"),
    ("Point focal BNEC", "Coordonne les validations et l'intégration.", "Validation, BNEC, rapports"),
    ("Cellule de veille", "Suit les échéances, alertes et relances.", "Veille, alertes, échéances"),
    ("Administrateur HAUQE", "Gère comptes, référentiels, règles et continuité.", "Administration, audit, sauvegardes")
], [3.2, 7.2, 6.1])
add_heading(doc, "Table des matières", 1)
p = doc.add_paragraph(); add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
add_body(doc, "Dans Word : clic droit sur la table des matières, puis « Mettre à jour les champs » avant impression.")
add_heading(doc, "Liste des figures", 2); add_body(doc, "Les figures seront mises à jour après insertion des captures d'écran.")
add_heading(doc, "Glossaire essentiel", 2)
add_kv_table(doc, [("BNEC", "Base nationale des entreprises certifiées."), ("FUCCS", "Grille de contrôle appliquée aux dossiers de certification."), ("INFC", "Indicateur de fiabilité des certifications."), ("SNCC", "Système national de classement des certifications."), ("CVC", "Cellule de veille des certifications.")])

add_chapter(doc, 1, "Comprendre le SNGSC et le cycle d'un dossier", "Situer chaque action dans le parcours national : aucune étape ne remplace une autre et chaque décision doit être traçable.", "Tous les utilisateurs habilités.", "Être connecté avec un compte actif.", [
    ("Lire le parcours d'un dossier", ["Rechercher l'entreprise ou la certification concernée.", "Ouvrir le dossier et consulter son statut courant.", "Lire l'historique avant d'effectuer une action.", "Utiliser uniquement l'action autorisée par votre rôle."], "Le dossier est traité dans l'ordre prévu : collecte, vérification, contrôle, validation, intégration, évaluation ou veille."),
    ("Identifier une action bloquante", ["Repérer les messages d'anomalie ou les statuts d'attente.", "Consulter la correction demandée ou le responsable affecté.", "Ne pas créer un second dossier pour contourner un blocage.", "Corriger ou transmettre au rôle compétent."], "Le prochain responsable peut reprendre le même dossier avec l'historique complet.")
], ["Un brouillon n'est pas une validation.", "Une intégration BNEC ne doit être ouverte qu'après les validations requises.", "Toute correction importante doit rester visible dans l'historique."], "Parcours global d'une certification", "Identifier les étapes, le statut courant, les actions autorisées et le lien vers l'historique.")

add_chapter(doc, 2, "Connexion, sécurité et verrouillage de session", "Accéder au système de manière sécurisée et protéger sa session de travail.", "Tous les utilisateurs disposant d'un compte actif.", "Adresse professionnelle, mot de passe et, si activée, application MFA.", [
    ("Se connecter", ["Ouvrir la page de connexion.", "Saisir l'adresse professionnelle et le mot de passe.", "Saisir le code temporaire MFA lorsqu'il est demandé.", "Vérifier que le tableau de bord et votre nom s'affichent."], "La session est ouverte avec les droits associés à votre rôle."),
    ("Récupérer l'accès", ["Choisir « Mot de passe oublié ».", "Saisir votre adresse professionnelle.", "Ouvrir le lien temporaire reçu par courriel.", "Définir un nouveau mot de passe et reconnectez-vous."], "Le lien est à usage unique ; les anciennes sessions peuvent être révoquées."),
    ("Déverrouiller une session", ["Saisir le code privé configuré lorsque l'écran de verrouillage apparaît.", "Vérifier l'avatar ou les initiales affichés.", "Après déverrouillage, reprendre le travail sur la même page."], "La session reprend sans perdre le contexte de travail." )
], ["Ne jamais communiquer un code MFA ou un code de récupération.", "Après plusieurs codes de verrouillage erronés, une déconnexion peut être imposée.", "Se déconnecter manuellement en quittant un poste partagé."], "Connexion et verrouillage", "Repérer les champs de connexion, l'accès MFA, le bouton mot de passe oublié et le verrouillage de session.")

add_chapter(doc, 3, "Profil, préférences et notifications", "Maintenir son identité professionnelle et configurer les préférences de travail sans modifier ses propres permissions.", "Tous les utilisateurs habilités.", "Être connecté ; disposer des informations personnelles à jour.", [
    ("Mettre à jour son profil", ["Ouvrir le menu utilisateur puis « Mon profil ».", "Corriger les coordonnées autorisées.", "Enregistrer et vérifier le message de confirmation.", "Ajouter une photo si nécessaire ; l'application conserve les initiales en solution de secours."], "Le profil, la navbar et l'écran de verrouillage présentent une identité cohérente."),
    ("Régler l'actualisation automatique", ["Ouvrir les préférences de profil.", "Activer ou désactiver l'actualisation automatique.", "Choisir un intervalle proposé : 15, 30, 60, 120 ou 300 secondes.", "Utiliser le bouton Actualiser pour un rechargement immédiat."], "L'utilisateur maîtrise le rafraîchissement de son écran sans perdre le bouton de contrôle manuel."),
    ("Traiter les notifications", ["Ouvrir la cloche de la navbar.", "Lire la notification et ouvrir la ressource concernée.", "Marquer la notification comme lue après traitement.", "Utiliser « Tout marquer comme lu » seulement après vérification."], "Les messages non lus correspondent aux traitements réellement à suivre." )
], ["Les préférences ne donnent pas de nouveaux droits.", "Une photo de profil doit rester professionnelle et ne contenir aucune donnée sensible.", "Une actualisation peut remplacer des modifications non enregistrées : enregistrer avant de rafraîchir."], "Profil et préférences", "Repérer la photo, les onglets de sécurité, les paramètres de rafraîchissement et les notifications.")

add_chapter(doc, 4, "Tableau de bord, alertes et échéances", "Prioriser le travail quotidien à partir des indicateurs, alertes et rendez-vous de suivi.", "Tous les rôles selon leurs droits de consultation et de traitement.", "Être connecté ; disposer d'au moins une donnée visible dans le périmètre utilisateur.", [
    ("Lire le tableau de bord", ["Observer les indicateurs et leur période de référence.", "Utiliser les filtres disponibles avant d'interpréter un résultat.", "Ouvrir une carte ou une ligne lorsque l'action doit être traitée.", "Actualiser si une mise à jour récente est attendue."], "Les priorités et leur origine sont identifiées avant intervention."),
    ("Traiter une alerte", ["Ouvrir la liste des alertes.", "Lire le titre, la source, la priorité et le délai.", "Affecter un responsable si votre rôle le permet.", "Ouvrir la ressource liée, effectuer l'action, puis résoudre l'alerte avec une trace."], "L'alerte est soit traitée, soit affectée, soit maintenue ouverte avec un responsable identifié."),
    ("Gérer une échéance", ["Ouvrir le calendrier ou la liste globale des échéances.", "Sélectionner l'échéance pour consulter ses détails.", "Terminer ou annuler l'échéance en motivant la décision.", "Vérifier que la mention d'exécution apparaît sur le calendrier."], "L'action est journalisée et les rappels restent cohérents avec le statut du dossier." )
], ["Les seuils usuels de veille sont 180, 90, 30 jours et l'expiration.", "Ne clôturez pas une échéance de renouvellement sans justificatif ou décision motivée.", "Une alerte ne remplace pas la consultation du dossier source."], "Alertes et échéances", "Repérer les filtres, le calendrier, les priorités, la ressource source et les boutons de traitement.")

add_chapter(doc, 5, "Registre national : entreprises, organismes, zones et certifications", "Consulter et maintenir les données de référence sans créer de doublons.", "Agents habilités du registre, vérificateurs, points focaux et administrateurs.", "Rechercher l'existant avant toute création.", [
    ("Créer ou compléter une entreprise", ["Rechercher par raison sociale, identifiant, RCCM ou NIF.", "Ouvrir l'entreprise trouvée ou choisir la création si aucun résultat fiable n'existe.", "Saisir les informations justifiées et associer la zone administrative.", "Enregistrer puis consulter l'historique de l'entreprise."], "Une seule fiche entreprise fiable est maintenue dans le registre."),
    ("Gérer une certification", ["Depuis l'entreprise, consulter les certifications liées.", "Créer la certification avec le référentiel, l'organisme, les dates et la portée.", "Joindre ou associer les justificatifs disponibles.", "Utiliser le parcours de renouvellement lorsque la certification est reconduite."], "La certification est rattachée au bon titulaire et ses échéances peuvent être suivies."),
    ("Précréer un organisme", ["Lorsqu'un organisme certificateur est absent, utiliser la précréation autorisée.", "Renseigner le nom connu et les coordonnées disponibles.", "Conserver le statut incomplet jusqu'à vérification.", "Compléter ultérieurement la fiche depuis le registre."], "La collecte peut avancer sans inventer de données définitives." )
], ["Une création incomplète doit être complétée, non oubliée.", "Les dates de début et d'expiration alimentent les échéances lorsqu'elles sont validées.", "Éviter les doublons par recherche préalable et rapprochement."], "Fiche entreprise et certification", "Repérer les onglets, le statut, l'historique, les documents, les certifications et les actions de modification.")

add_chapter(doc, 6, "Collecte : campagnes, missions, déclarations et preuves", "Constituer une fiche de collecte complète et soumettre un dossier exploitable aux étapes suivantes.", "Agent de collecte, superviseur ou point focal selon les permissions.", "Campagne ou mission disponible ; entreprise existante ou précréation autorisée.", [
    ("Préparer une mission", ["Sélectionner une campagne existante ou créer la campagne si votre rôle l'autorise.", "Renseigner la zone, les dates, la priorité et l'objet de la mission.", "Affecter l'agent lorsque l'autorisation COLLECTE.AFFECTER est disponible.", "Enregistrer avant de passer à l'étape entreprise."], "La mission porte une référence, une zone, des dates et une responsabilité claire."),
    ("Saisir l'entreprise et le déclarant", ["Rechercher l'entreprise dans le registre.", "Précréer une entreprise seulement si aucun résultat n'est fiable.", "Renseigner le déclarant, sa fonction, ses contacts et le consentement.", "Conserver une mention ou référence de signature lorsque le consentement est recueilli."], "La fiche identifie le titulaire et l'interlocuteur ayant fourni les informations."),
    ("Déclarer les produits et certifications", ["Ajouter chaque produit ou service séparément avec catégorie, volume, unité et marchés visés.", "Ajouter les certifications déclarées avec numéro, organisme, norme, dates et portée.", "Choisir la situation actuelle : présente, absente, audit de surveillance, renouvellement, expirée ou audit initial.", "Ne pas transformer cette déclaration terrain en certification BNEC officielle."], "Les déclarations sont prêtes pour rapprochement et vérification."),
    ("Déposer les preuves et soumettre", ["Sélectionner les PDF ou images acceptés.", "Vérifier que chaque fichier apparaît comme sélectionné avant envoi.", "Saisir les observations utiles, puis consulter le contrôle final.", "Soumettre la fiche seulement lorsque les informations essentielles sont cohérentes."], "La fiche passe au statut de traitement suivant et les pièces restent rattachées au dossier." )
], ["Le bouton Continuer crée la mission et met à jour l'URL de reprise : ne le relancez pas pour créer une seconde mission.", "Une déclaration de certificat expiré doit être justifiée par les dates et la preuve disponible.", "Enregistrer un brouillon avant de quitter une saisie longue."], "Formulaire de collecte", "Repérer l'étape courante, la progression, les champs obligatoires, la sélection de fichiers et le bouton Continuer ou Soumettre.")

add_chapter(doc, 7, "Vérification documentaire et contrôle FUCCS", "Confirmer la cohérence des informations collectées et documenter les anomalies avant validation.", "Vérificateur, contrôleur FUCCS, superviseur et point focal autorisé.", "Fiche soumise et dossier affecté ; pièces consultables.", [
    ("Vérifier un dossier", ["Ouvrir le dossier affecté.", "Comparer les informations de la fiche aux pièces jointes et aux sources disponibles.", "Confirmer les informations cohérentes ou signaler une anomalie motivée.", "Renseigner les dates de traitement lorsque la planification est requise."], "Le dossier conserve une position claire : confirmé, à corriger ou à compléter."),
    ("Conduire le contrôle FUCCS", ["Ouvrir la grille associée au dossier.", "Démarrer le contrôle une seule fois ; utiliser Reprendre pour mettre à jour le même résultat.", "Renseigner les critères, notes, constats et preuves.", "Finaliser lorsque tous les critères requis sont renseignés."], "Un résultat FUCCS unique et historisé est disponible pour le dossier."),
    ("Réouvrir ou corriger", ["Utiliser Réouvrir uniquement lorsqu'une correction justifiée est nécessaire.", "Documenter le motif et l'impact sur le dossier.", "Mettre à jour l'ancien contrôle au lieu d'en créer un nouveau.", "Vérifier l'historique après l'action."], "La traçabilité montre la décision, le motif et la nouvelle version de travail." )
], ["Ne pas créer plusieurs contrôles FUCCS pour le même dossier.", "Une anomalie doit être compréhensible et accompagnée d'une action attendue.", "Les dates d'affectation peuvent alimenter les alertes et échéances."], "Dossier de vérification et grille FUCCS", "Repérer le responsable, les dates, les critères, les preuves, les constats et les boutons Démarrer, Reprendre ou Réouvrir.")

add_chapter(doc, 8, "Validation N1/N2 et intégration BNEC", "Prendre une décision de validation fondée et intégrer uniquement les dossiers prêts dans le registre officiel.", "Valideurs N1/N2, Direction Technique, point focal BNEC et administrateur fonctionnel autorisé.", "Vérification et contrôle disponibles ; corrections traitées ou justifiées.", [
    ("Réaliser une validation", ["Ouvrir le dossier de validation et lire les conclusions précédentes.", "Contrôler les pièces, les constats et les corrections apportées.", "Émettre la décision N1 ou N2 avec un motif clair.", "Renvoyer le dossier en correction lorsque les éléments restent insuffisants."], "La décision est explicite, motivée et inscrite dans l'historique."),
    ("Préparer l'intégration BNEC", ["Ouvrir l'espace d'intégration seulement après la validation requise.", "Vérifier le plan d'intégration et les éléments bloquants.", "Contrôler la codification, le titulaire, l'organisme, le référentiel et les dates.", "Lancer l'intégration puis consulter le post-contrôle."], "Le dossier est intégré une seule fois dans la BNEC, ou son échec est documenté."),
    ("Traiter un échec d'intégration", ["Lire le motif d'échec et l'élément concerné.", "Corriger le dossier source ou transmettre au rôle compétent.", "Utiliser Nouvelle tentative après correction.", "Ne pas contourner le post-contrôle ou les vérifications obligatoires."], "La nouvelle tentative reste liée à l'historique de l'intégration initiale." )
], ["Une validation favorable ne remplace pas les contrôles de préparation BNEC.", "La codification ne doit pas être saisie au hasard : utiliser les règles et référentiels publiés.", "L'intégration prépare les échéances, audits et alertes liés à la certification."], "Validation et intégration BNEC", "Repérer les décisions N1/N2, les corrections, le plan d'intégration, les blocages et le post-contrôle.")

add_chapter(doc, 9, "Scoring, INFC et classement SNCC", "Évaluer les données sans dupliquer les résultats et comprendre la portée de chaque indicateur.", "Utilisateurs autorisés au scoring, à l'INFC et au SNCC.", "Modèle publié et données préalables disponibles ; INFC validé avant SNCC lorsque la règle l'exige.", [
    ("Évaluer une entreprise", ["Ouvrir la liste des entreprises éligibles au scoring.", "Sélectionner le modèle actif et renseigner les valeurs demandées.", "Valider l'évaluation selon les droits disponibles.", "Cliquer sur la ligne historique pour consulter les évaluations précédentes."], "Une seule ligne courante est conservée par entreprise ; ses évaluations successives restent consultables."),
    ("Calculer l'INFC", ["Ouvrir les certifications éligibles.", "Vérifier le modèle, la formule et les informations d'entrée.", "Lancer le calcul ou le recalcul de la certification.", "Consulter la ligne courante et son historique de recalculs."], "Le dernier résultat remplace la ligne courante sans supprimer les calculs antérieurs."),
    ("Classer au SNCC", ["Sélectionner une certification dont l'INFC est validé si la règle l'impose.", "Choisir les valeurs de classement autorisées et justifier la décision.", "Valider le classement, ou clôturer la période avec un motif.", "Cliquer sur la ligne pour consulter l'évolution des classements."], "Le classement courant est lisible et les décisions antérieures sont historisées." )
], ["Un modèle sans niveaux institutionnels peut permettre le calcul mais bloquer la validation finale.", "Ne créez pas cinquante lignes pour une même entité : utilisez recalcul, réévaluation ou reclassement.", "Consultez la légende SNCC avant de valider une classe ou un niveau de risque."], "Résultats de scoring, INFC et SNCC", "Repérer le modèle actif, la valeur courante, la date, le statut, la ligne cliquable et le modal d'historique.")

add_chapter(doc, 10, "Veille, relances, décisions et communication", "Suivre les risques, échéances et réponses externes jusqu'à leur clôture motivée.", "Cellule de veille, superviseurs, agents affectés et décideurs autorisés.", "Dossier de veille ou échéance existante ; destinataire et responsable identifiés.", [
    ("Ouvrir un dossier de veille", ["Créer le dossier à partir d'un risque, d'une certification, d'une entreprise ou d'une alerte.", "Choisir la nature de suivi, la priorité, le responsable et la prochaine action.", "Renseigner la date de prochaine action afin de créer le rappel.", "Ajouter les événements et pièces de suivi."], "Le dossier est affecté, visible dans la file de veille et associé à ses rappels."),
    ("Programmer une relance", ["Saisir le destinataire, son adresse électronique, le contenu du message et la date d'envoi.", "Renseigner également le délai de réponse.", "Enregistrer : la date d'envoi et l'échéance sont placées dans le calendrier.", "Si la date d'envoi est aujourd'hui, vérifier l'envoi dans la file de notifications."], "La relance est journalisée et une alerte est générée à l'approche ou au dépassement du délai."),
    ("Créer une décision", ["Ouvrir Décisions et actions.", "Choisir l'objet, le contexte, l'autorité et les responsables concernés.", "Utiliser les propositions de saisie puis adapter le contenu au cas réel.", "Enregistrer et vérifier la notification des utilisateurs concernés."], "La décision est traçable et les actions associées peuvent être suivies." )
], ["Vérifier l'adresse électronique avant une programmation d'envoi.", "Une clôture de dossier ou d'échéance doit toujours être motivée.", "Ne pas inscrire de données confidentielles inutiles dans le corps d'un courriel."], "Dossier de veille et nouvelle relance", "Repérer la file de veille, le responsable, les dates, le destinataire, le message et les actions de clôture.")

add_chapter(doc, 11, "Administration, référentiels et publication", "Administrer les paramètres métier et les utilisateurs sans modifier les données opérationnelles sans justification.", "Administrateur HAUQE, administrateur fonctionnel et rôles expressément autorisés.", "Permission d'administration adaptée ; référentiel ou règle de travail identifié.", [
    ("Créer un utilisateur", ["Ouvrir Administration puis Utilisateurs.", "Renseigner l'identité professionnelle, l'adresse et les rôles initiaux.", "Choisir l'envoi des identifiants seulement lorsque la procédure l'autorise.", "Enregistrer et vérifier le statut actif du compte."], "Le compte reçoit uniquement les habilitations prévues par les rôles sélectionnés."),
    ("Gérer un référentiel ou une règle", ["Rechercher d'abord le code ou la valeur existante.", "Créer un brouillon avec un code unique proposé ou saisi selon les règles.", "Ajouter les valeurs, descriptions et conditions nécessaires.", "Soumettre ou publier uniquement après la validation institutionnelle requise."], "Les règles et valeurs publiées deviennent des références traçables pour l'application."),
    ("Préparer une publication", ["Produire les données nominatives ou agrégées depuis les règles prévues.", "Créer une demande de publication au statut brouillon.", "Soumettre à approbation, puis publier après décision.", "Retirer une publication avec motif si son maintien n'est plus justifié."], "Le tableau public est alimenté uniquement par des données autorisées et validées." )
], ["Les rôles sont cumulatifs : vérifier l'impact avant une attribution.", "Ne publier aucune donnée nominative sans validation et respect du périmètre autorisé.", "Une valeur de référentiel modifiée peut influencer plusieurs pages : documenter le changement."], "Administration et publication", "Repérer les onglets utilisateurs, rôles, référentiels, règles, brouillons, approbations et publications retirées.")

add_chapter(doc, 12, "Rapports, journal d'audit, qualité et sauvegardes", "Produire des informations fiables, vérifier la traçabilité et protéger la continuité du système.", "Utilisateurs autorisés aux rapports, à l'audit, à la qualité et aux sauvegardes.", "Droits appropriés ; période, périmètre et objectif de l'opération définis.", [
    ("Générer un rapport", ["Ouvrir Rapports et choisir le modèle approprié.", "Définir la période et les filtres nécessaires.", "Choisir le format PDF, Excel ou CSV.", "Contrôler l'en-tête HAUQE, les tableaux, les dates et le périmètre avant diffusion."], "Le rapport est conservé dans l'historique avec son format, sa date et son demandeur."),
    ("Consulter le journal d'audit", ["Ouvrir Journal d'audit.", "Filtrer par utilisateur, période, type d'action ou ressource.", "Lire les valeurs avant/après lorsque disponibles.", "Exporter uniquement les informations nécessaires à l'analyse."], "Les événements sont compréhensibles et le nom des utilisateurs autorisés reste lisible."),
    ("Lancer ou suivre une sauvegarde", ["Ouvrir Sauvegardes et choisir système, documents ou complète.", "Pour une sauvegarde manuelle, confirmer l'action et attendre la progression.", "Vérifier le résultat, l'empreinte ou le statut final.", "Planifier les sauvegardes automatiques et effectuer périodiquement un test de restauration supervisé."], "Une archive cohérente est produite, sans restaurer directement en production." )
], ["Une sauvegarde manuelle peut bloquer temporairement certaines actions : prévenir les utilisateurs concernés.", "Ne pas télécharger ni partager une sauvegarde hors du périmètre autorisé.", "Le journal d'audit sert à comprendre une action ; il ne remplace pas une enquête ou une décision formelle."], "Rapports, audit et sauvegardes", "Repérer le choix de rapport, les filtres, l'historique, les types de sauvegarde, la progression et le résultat final.")

add_page_break(doc)
add_heading(doc, "Annexes", 1)
add_heading(doc, "Checklist avant une action sensible", 2)
add_bullets(doc, ["J'ai ouvert le bon dossier et lu son historique.", "Je dispose de la permission nécessaire.", "Les données saisies sont appuyées par une pièce ou une source identifiable.", "La décision ou la clôture est motivée.", "Je sais quel utilisateur ou quel module recevra la suite du traitement.", "Je n'ai pas exposé de données sensibles dans un commentaire, une capture ou un courriel."])
add_heading(doc, "Erreurs fréquentes et première réaction", 2)
add_matrix(doc, ["Situation", "Réaction attendue"], [
    ("Bouton sans réaction", "Attendre la fin de l'indicateur de chargement, actualiser si nécessaire, puis signaler l'erreur avec la page et l'heure."),
    ("Accès refusé", "Vérifier le rôle attribué ; ne pas demander le mot de passe d'un autre utilisateur."),
    ("Doublon détecté", "Rechercher l'entité existante et utiliser le dossier courant ou la fonction de rapprochement."),
    ("Courriel non reçu", "Vérifier l'adresse, la file de notification et le délai prévu avant de relancer."),
    ("Donnée incohérente", "Ne pas l'écraser sans preuve ; utiliser correction, anomalie ou décision selon le parcours."),
], [4.6, 11.9])
add_heading(doc, "Historique des versions du guide", 2)
add_matrix(doc, ["Version", "Date", "Évolution"], [("1.0", "Août 2026", "Refonte opérationnelle : procédures détaillées, rôles, contrôles et espaces de captures.")], [3, 3.2, 10.3])

doc.core_properties.title = "Guide global d'utilisation SNGSC / HAUQE"
doc.core_properties.subject = "Manuel opérationnel des utilisateurs habilités"
doc.core_properties.author = "HAUQE"
doc.save(OUT)
print(OUT)
