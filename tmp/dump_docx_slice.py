import sys
from pathlib import Path
from docx import Document


path = Path(sys.argv[1])
start = int(sys.argv[2])
end = int(sys.argv[3])
doc = Document(path)
for idx in range(start, min(end, len(doc.paragraphs))):
    paragraph = doc.paragraphs[idx]
    print(f"P{idx:03d} [{paragraph.style.name}] {paragraph.text}")
