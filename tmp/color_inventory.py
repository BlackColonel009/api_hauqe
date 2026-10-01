from collections import Counter, defaultdict
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def iter_paragraphs(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


for name in [
    "Guide_global_utilisation_SNGSC_HAUQE_v2_mis_a_jour_final.docx",
    "Guide_global_utilisation_SNGSC_HAUQE_v3.docx",
]:
    doc = Document(ROOT / "output" / "docx" / name)
    total = Counter()
    by_style = defaultdict(Counter)
    for paragraph in iter_paragraphs(doc):
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            rgb = run.font.color.rgb
            key = str(rgb) if rgb else "AUTO"
            total[key] += 1
            by_style[paragraph.style.name][key] += 1
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in iter_paragraphs(part):
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    rgb = run.font.color.rgb
                    key = str(rgb) if rgb else "AUTO"
                    total[key] += 1
                    by_style[paragraph.style.name][key] += 1
    print(f"\n{name}\nTOTAL {total}")
    for style, colors in sorted(by_style.items()):
        print(f"{style}: {colors}")
