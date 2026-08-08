from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import build_guide_global_squelette as base


OUTPUT = Path("docs_md/guide_global_utilisation_SNGSC_HAUQE_version_redigee.docx")


MODULES = [
    {
        "title": "Comprendre le SNGSC et son parcours de traitement",
        "roles": "Tous les utilisateurs habilités, selon les droits attribués.",
        "purpose": "Situer chaque action dans le cycle national de gestion et de suivi des certifications.",
        "intro": "Le SNGSC centralise les données des entreprises, organismes certificateurs, certifications, collectes et contrôles. Il conserve l'historique des décisions, des pièces, des échanges et des actions de suivi afin de rendre le traitement traçable.",
        "steps": [
            ("Identifier le dossier", "Rechercher l'entreprise ou la certification dans le registre, puis ouvrir son dossier."),
            ("Suivre l'étape courante", "Lire le statut et l'historique avant toute modification."),
            ("Traiter la file appropriée", "Intervenir dans Collecte, Vérification, Contrôle, Validation, Intégration ou Veille selon l'étape affichée."),
        ],
        "attention": "Le cycle de référence est : Brouillon, Soumise, Vérification, Contrôle, Validation définitive, Intégration BNEC, classification/INFC, SNCC, puis Veille. Une étape ne remplace pas une autre.",
        "frames": ["Parcours global d'une certification"],
    },
    {
        "title": "Connexion, sécurité et verrouillage de session",
        "roles": "Tous les utilisateurs disposant d'un compte actif.",
        "purpose": "Accéder au système de manière sécurisée et protéger la session de travail.",
        "intro": "La connexion utilise l'adresse professionnelle et le mot de passe. Les mécanismes de sécurité, notamment la MFA lorsqu'elle est activée, le verrouillage et la récupération du mot de passe, dépendent des droits et de la configuration de l'environnement.",
        "steps": [
            ("Se connecter", "Saisir l'adresse professionnelle et le mot de passe, puis suivre l'étape MFA si elle est demandée."),
            ("Récupérer l'accès", "Utiliser « Mot de passe oublié » et suivre le lien temporaire reçu. Ne jamais partager ce lien."),
            ("Verrouiller la session", "Depuis Sécurité, définir le code de verrouillage et l'inactivité souhaitée, puis utiliser « Tester » si nécessaire."),
        ],
        "attention": "Après plusieurs tentatives de code erronées, la session peut être fermée. Un mot de passe ou un code de verrouillage ne doit jamais être transmis par courriel ou messagerie.",
        "frames": ["Écran de connexion", "Paramètres de verrouillage de session"],
    },
    {
        "title": "Profil, préférences et notifications",
        "roles": "Tous les utilisateurs habilités.",
        "purpose": "Maintenir ses coordonnées, ses préférences d'affichage et ses paramètres de notification.",
        "intro": "Le menu utilisateur donne accès au profil, aux préférences, à la sécurité, aux sessions actives et à la déconnexion. Les droits et rôles sont administrés séparément par les personnes habilitées.",
        "steps": [
            ("Mettre à jour le profil", "Ouvrir Profil et préférences, corriger les coordonnées autorisées puis enregistrer."),
            ("Régler l'actualisation", "Activer ou désactiver l'actualisation automatique, choisir l'intervalle, puis conserver le bouton « Actualiser » pour un contrôle immédiat."),
            ("Gérer les notifications", "Ouvrir la cloche, consulter les messages, les marquer comme lus et suivre le lien vers le module concerné."),
        ],
        "attention": "L'avatar est facultatif. L'application revient aux initiales si l'image n'est pas disponible. Les notifications reflètent le périmètre d'accès de l'utilisateur.",
        "frames": ["Profil et préférences", "Centre de notifications"],
    },
    {
        "title": "Tableau de bord, alertes et échéances",
        "roles": "Direction, responsables de suivi et profils de consultation autorisés.",
        "purpose": "Prioriser le travail quotidien à partir des indicateurs, alertes et échéances.",
        "intro": "Le tableau de bord donne une vue synthétique. Une carte, une ligne récente ou une action prioritaire ouvre le registre ou le dossier correspondant. Les données affichées doivent toujours être lues avec leur période, leurs filtres et leur statut de mise à jour.",
        "steps": [
            ("Lire les priorités", "Consulter d'abord les alertes critiques, les échéances proches et les tâches assignées."),
            ("Filtrer", "Choisir la période, le territoire, le type de certification ou le statut adapté au besoin."),
            ("Ouvrir le détail", "Cliquer sur une carte, une ligne ou l'action associée pour poursuivre dans le module concerné."),
        ],
        "attention": "Les horizons de référence sont 180 jours, 90 jours, 30 jours et l'expiration. Une alerte critique reste ouverte jusqu'à régularisation ou clôture motivée.",
        "frames": ["Tableau de bord national", "Centre des alertes", "Calendrier des échéances"],
    },
    {
        "title": "Registre national : entreprises, organismes, zones et certifications",
        "roles": "Agents habilités à consulter ou gérer les registres.",
        "purpose": "Consulter, créer ou mettre à jour les entités de référence sans perdre leur historique.",
        "intro": "Les pages Entreprises, Organismes, Zones administratives et Certifications fonctionnent comme des registres liés. Avant de créer un enregistrement, il convient de rechercher les homonymes, doublons et éléments déjà enregistrés.",
        "steps": [
            ("Rechercher", "Utiliser la recherche et les filtres avant de sélectionner « Nouveau ». Ouvrir le détail pour contrôler l'historique."),
            ("Créer ou corriger", "Saisir les données autorisées, joindre les pièces demandées et enregistrer. Pour une donnée déjà utilisée, préférer la mise à jour ou la désactivation à la suppression."),
            ("Vérifier les liens", "Contrôler les organismes, accréditations, portées, dates, zones et certifications reliés avant de valider la saisie."),
        ],
        "attention": "Une entreprise sans RCCM peut être conservée avec un statut de régularisation. Un organisme non accrédité peut être enregistré, mais les certificats associés doivent être signalés pour vérification.",
        "frames": ["Liste des entreprises", "Dossier d'une entreprise", "Registre des certifications", "Dossier d'un organisme"],
    },
    {
        "title": "Collecte : campagnes, missions, déclarations et preuves",
        "roles": "Responsables de campagne, agents de collecte et profils de supervision autorisés.",
        "purpose": "Préparer une mission, saisir la collecte numérique et soumettre un dossier complet.",
        "intro": "La collecte relie une campagne ou une mission, une entreprise, des activités, produits, marchés, certifications déclarées et justificatifs. Le brouillon permet de poursuivre le travail ultérieurement sans transmettre un dossier incomplet.",
        "steps": [
            ("Préparer la mission", "Choisir la campagne et la mission, puis vérifier l'entreprise et le périmètre de collecte."),
            ("Renseigner la fiche", "Compléter les étapes du formulaire, ajouter les produits, offres, certifications et situations déclarées nécessaires."),
            ("Joindre et soumettre", "Contrôler les fichiers sélectionnés, retirer ceux qui sont erronés, enregistrer un brouillon si besoin puis soumettre lorsque les contrôles de complétude sont satisfaits."),
        ],
        "attention": "Les justificatifs affichent leur nom et leur taille avant l'envoi. Les critères FUCCS et les décisions ne sont pas saisis pendant la collecte ; ils relèvent des étapes ultérieures.",
        "frames": ["Liste des collectes et missions", "Formulaire de collecte", "Étape Preuves et observations"],
    },
    {
        "title": "Vérification documentaire et confirmations externes",
        "roles": "Agents vérificateurs, Point focal BNEC et responsables de suivi autorisés.",
        "purpose": "Établir la recevabilité documentaire d'un dossier avant le contrôle approfondi.",
        "intro": "La vérification porte sur la complétude, les pièces, les dates, la portée, le référentiel, l'organisme certificateur et son accréditation. Elle permet également de consigner les demandes adressées aux organismes et les réponses reçues.",
        "steps": [
            ("Prendre en charge le dossier", "Ouvrir la file de vérification, prendre connaissance de l'affectation et du contexte du dossier."),
            ("Contrôler les éléments", "Examiner les pièces, les informations déclarées, les dates de validité et la cohérence avec l'organisme et l'accréditation."),
            ("Conclure et orienter", "Enregistrer l'avis, les réserves ou la demande de complément. Conserver toute confirmation externe et orienter le dossier vers l'étape suivante autorisée."),
        ],
        "attention": "Un avis de vérification peut être conforme, sous réserve, non vérifié, suspect ou rejeté selon les règles applicables. Une demande de complément doit être motivée et traçable.",
        "frames": ["File de vérification", "Détail d'un dossier de vérification", "Demande de confirmation externe"],
    },
    {
        "title": "Contrôle FUCCS et constats",
        "roles": "Agents chargés du contrôle FUCCS et responsables autorisés.",
        "purpose": "Évaluer les critères de contrôle, documenter les constats et finaliser le résultat du contrôle.",
        "intro": "La grille FUCCS est versionnée. La version active visible à l'écran peut évoluer ; il faut donc utiliser les critères effectivement chargés pour le dossier et ne pas présumer d'un nombre fixe de critères. Les notes, preuves et constats restent associés au contrôle concerné.",
        "steps": [
            ("Ouvrir le contrôle", "Depuis la file ou le dossier, sélectionner le contrôle à traiter et vérifier la version de grille appliquée."),
            ("Noter et justifier", "Attribuer une note selon les règles affichées, joindre ou référencer les preuves et enregistrer les constats nécessaires."),
            ("Finaliser", "Relire les notes et constats, puis finaliser le contrôle uniquement lorsque les éléments obligatoires sont renseignés."),
        ],
        "attention": "Le score FUCCS, l'INFC et le classement SNCC sont trois résultats distincts. Toute réouverture ou correction doit préserver l'historique et le motif.",
        "frames": ["Espace de travail FUCCS", "Saisie d'une note et d'un constat", "Synthèse du contrôle"],
    },
    {
        "title": "Validation N1/N2 et décisions de dossier",
        "roles": "Responsables de validation habilités et Direction Technique selon la matrice de droits.",
        "purpose": "Enregistrer une décision hiérarchisée après vérification et contrôle, avec ses motifs et réserves.",
        "intro": "La validation est distincte de la vérification documentaire et du contrôle FUCCS. Le dossier doit être traité dans la file correspondant à son statut ; les décisions, réserves et justifications deviennent des éléments de l'historique.",
        "steps": [
            ("Examiner le dossier", "Consulter les résultats de vérification, le contrôle et les pièces avant de préparer la décision."),
            ("Choisir la décision", "Sélectionner la décision autorisée, formuler les motifs, les réserves éventuelles et les demandes de correction."),
            ("Enregistrer la suite", "Confirmer la décision. Un dossier définitivement validé devient éligible à l'intégration BNEC ; un retour ou un rejet reste documenté."),
        ],
        "attention": "Ne pas utiliser la validation pour contourner une étape manquante. Les niveaux, visas et règles de décision applicables doivent être paramétrés par la HAUQE.",
        "frames": ["File des validations", "Détail d'une validation", "Formulaire de décision"],
    },
    {
        "title": "Intégration BNEC et mise à jour du registre officiel",
        "roles": "Administrateurs fonctionnels, Point focal BNEC et responsables habilités.",
        "purpose": "Contrôler les données validées, éviter les doublons et tracer leur intégration dans la BNEC.",
        "intro": "L'intégration BNEC intervient après une validation définitive. L'espace de travail présente la file, les données candidates, les codes proposés et les contrôles préalables. Chaque action doit rester reliée au dossier source.",
        "steps": [
            ("Sélectionner une entrée éligible", "Ouvrir la file d'intégration et vérifier l'état de validation ainsi que les données proposées."),
            ("Effectuer les contrôles", "Rechercher les doublons, contrôler la codification, corriger les écarts autorisés et consigner les anomalies."),
            ("Intégrer et contrôler après coup", "Démarrer l'intégration, vérifier le résultat, réaliser le contrôle post-intégration puis clôturer l'opération."),
        ],
        "attention": "Une intégration ne doit jamais supprimer l'historique de la collecte ou de la décision. Le recalcul des indicateurs et la traçabilité des modifications suivent les règles publiées.",
        "frames": ["File d'intégration BNEC", "Détail d'une intégration", "Contrôle post-intégration"],
    },
    {
        "title": "Scoring, INFC et classement SNCC",
        "roles": "Profils d'analyse, responsables métiers et décideurs autorisés.",
        "purpose": "Lire et utiliser correctement les trois résultats calculés par le système.",
        "intro": "Le scoring de contrôle, l'Indice national de fiabilité des certifications (INFC) et le classement SNCC répondent à des objectifs différents. L'INFC est construit à partir d'un modèle versionné, avec six domaines pondérés : authenticité, validité, maintien, maîtrise documentaire, traçabilité et maîtrise opérationnelle, suivi et renouvellement.",
        "steps": [
            ("Choisir le bon résultat", "Pour la conformité du contrôle, consulter FUCCS ; pour la fiabilité d'une certification, consulter l'INFC ; pour la situation de classement, consulter le SNCC."),
            ("Lire le contexte", "Vérifier le dossier, la date, la version de modèle et les données sources avant d'interpréter la valeur."),
            ("Déclencher le suivi", "Depuis les résultats ou actions prioritaires, créer une échéance ou un plan d'action lorsque les droits le permettent."),
        ],
        "attention": "Les pondérations, seuils, classes et règles de reclassement sont administrés et versionnés. Ils ne doivent pas être modifiés à partir d'une interprétation manuelle d'un résultat.",
        "frames": ["Classification entreprise et scoring", "Calcul INFC", "Classement SNCC"],
    },
    {
        "title": "Veille, relances, décisions et plans d'action",
        "roles": "Cellule de Veille des Certifications, Direction Technique et responsables désignés.",
        "purpose": "Suivre les échéances, relancer les parties concernées et piloter les actions décidées.",
        "intro": "La Cellule de Veille des Certifications organise les dossiers à surveiller, les relances, les réponses, les notes mensuelles et les rapports trimestriels. Les décisions et plans d'action relient un constat à une mesure, un responsable, une échéance et un indicateur de suivi.",
        "steps": [
            ("Ouvrir la file de veille", "Filtrer les dossiers par priorité, échéance, statut ou entité concernée."),
            ("Enregistrer une relance", "Choisir le canal, le destinataire, le délai attendu et joindre les éléments nécessaires. Enregistrer la réponse dès réception."),
            ("Suivre l'action", "Créer ou mettre à jour le plan d'action, renseigner l'avancement et évaluer l'effet avant clôture."),
        ],
        "attention": "Une relance, une décision ou une action clôturée doit conserver son motif, ses pièces et sa date. L'absence de réponse ne doit pas être effacée : elle fait partie du suivi.",
        "frames": ["File de veille", "Relance et échange", "Plan d'action"],
    },
    {
        "title": "Tableaux tactique, stratégique, annuel, national et public",
        "roles": "Direction, responsables de pilotage et profils de consultation autorisés.",
        "purpose": "Produire une lecture adaptée au niveau de pilotage sans divulguer de données sensibles.",
        "intro": "Les tableaux de bord sont organisés par finalité : opérationnelle, tactique, stratégique, annuelle, nationale ou publique. Le tableau public ne doit présenter que des données agrégées, non confidentielles et publiées selon le circuit autorisé.",
        "steps": [
            ("Choisir le niveau", "Sélectionner la vue correspondant à la décision à prendre : suivi quotidien, gestion, direction, bilan annuel ou publication."),
            ("Appliquer les filtres", "Définir la période, le territoire et le périmètre avant de lire les indicateurs."),
            ("Contrôler avant diffusion", "Pour toute diffusion, vérifier le caractère agrégé, la période, les sources et les autorisations de publication."),
        ],
        "attention": "Un indicateur n'est pas une décision. Avant de comparer deux périodes, vérifier la même définition, la même source et le même périmètre.",
        "frames": ["Tableau tactique", "Tableau stratégique", "Vue publique agrégée"],
    },
    {
        "title": "Administration : utilisateurs, référentiels, règles et publications",
        "roles": "Administrateurs fonctionnels et administrateurs système, selon les droits attribués.",
        "purpose": "Administrer les paramètres fonctionnels tout en garantissant la traçabilité des changements.",
        "intro": "L'administration couvre les utilisateurs, rôles, permissions, référentiels, règles métier, modèles de codification, documents et publications. Les opérations sensibles sont soumises aux droits et doivent être journalisées ; les mots de passe ne sont jamais visibles par les administrateurs.",
        "steps": [
            ("Gérer un utilisateur", "Créer ou rechercher le compte, renseigner les informations professionnelles, attribuer les droits nécessaires et utiliser le parcours sécurisé d'invitation ou de réinitialisation."),
            ("Mettre à jour un référentiel", "Rechercher l'élément, évaluer ses dépendances puis créer, modifier ou désactiver la valeur selon la procédure autorisée."),
            ("Publier une règle ou une donnée", "Préparer un brouillon, contrôler le contenu, renseigner le motif et la référence d'autorisation, puis appliquer le circuit de publication prévu."),
        ],
        "attention": "Une règle, une formule ou une grille publiée doit rester reproductible pour les dossiers historiques. Préférer une nouvelle version à l'écrasement d'une version utilisée.",
        "frames": ["Utilisateurs et accès", "Référentiels", "Règles et codification", "Publications"],
    },
    {
        "title": "Audit, qualité des données, sauvegarde et restauration",
        "roles": "Profils habilités à l'audit, à la qualité ou à l'administration système.",
        "purpose": "Contrôler la qualité, suivre les incidents et assurer la continuité des données.",
        "intro": "Le journal d'audit est en lecture seule. Il aide à comprendre qui a réalisé une opération, quand et dans quel contexte, sans exposer de secret. Les revues de qualité, incidents, archives et sauvegardes sont des processus distincts mais complémentaires.",
        "steps": [
            ("Rechercher une trace", "Dans le journal d'audit, filtrer par période, utilisateur, action ou objet, puis consulter le contexte disponible."),
            ("Traiter une anomalie", "Créer ou mettre à jour l'incident ou l'action d'amélioration, affecter un responsable et suivre la résolution."),
            ("Contrôler la continuité", "Consulter les sauvegardes, leur statut et les preuves de restauration. Une restauration doit être réalisée selon une procédure autorisée."),
        ],
        "attention": "Les journaux ne doivent contenir ni mot de passe, ni empreinte de mot de passe, ni jeton d'accès. La suppression définitive de données métier n'est pas le mode normal de traitement.",
        "frames": ["Journal d'audit", "Qualité des données", "Sauvegardes et restaurations"],
    },
    {
        "title": "Rapports, exportations, assistance et erreurs fréquentes",
        "roles": "Utilisateurs disposant des permissions de consultation ou d'export nécessaires.",
        "purpose": "Générer des documents exploitables, sécuriser les exports et résoudre les difficultés courantes.",
        "intro": "Le module Rapports permet de choisir un modèle, un périmètre, des sections et un format. Les exports restent soumis aux droits, à un motif lorsque celui-ci est demandé et au respect de la confidentialité des données.",
        "steps": [
            ("Préparer le rapport", "Choisir le modèle, la période, le périmètre et les sections à inclure."),
            ("Contrôler l'aperçu", "Vérifier le titre, les filtres, les indicateurs et l'absence de données non autorisées avant de générer."),
            ("Générer et archiver", "Choisir PDF, Excel ou CSV selon l'usage, récupérer le fichier dans l'historique puis le conserver selon les règles de diffusion."),
        ],
        "attention": "Si une action semble indisponible, vérifier d'abord les filtres, le statut du dossier et les droits attribués. Ne pas partager un export contenant des données nominatives sans autorisation.",
        "frames": ["Catalogue de rapports", "Paramétrage d'un export", "Historique des générations"],
    },
]


def add_text(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.add_run(text)


def add_procedure(doc, steps):
    doc.add_paragraph("Procédure", style="Heading 2")
    table = doc.add_table(rows=1, cols=3)
    base.set_table_geometry(table, [850, 2580, 5930])
    for cell, label in zip(table.rows[0].cells, ["Étape", "Action", "Comment procéder"]):
        base.set_cell_shading(cell, base.GREEN)
        base.set_cell_border(cell, base.GREEN_DARK)
        base.set_font(cell.paragraphs[0].add_run(label), 9, base.WHITE, bold=True)
    base.set_repeat_table_header(table.rows[0])
    for index, (action, instruction) in enumerate(steps, start=1):
        cells = table.add_row().cells
        for cell in cells:
            base.set_cell_border(cell)
        base.set_font(cells[0].paragraphs[0].add_run(str(index)), 9.5, base.GREEN_DARK, bold=True)
        base.set_font(cells[1].paragraphs[0].add_run(action), 9.5, base.GREEN_DARK, bold=True)
        base.set_font(cells[2].paragraphs[0].add_run(instruction), 9.5, "222222")


def add_frame(doc, figure_no, subject):
    base.add_capture_frame(doc, figure_no)
    paragraph = doc.paragraphs[-1]
    paragraph.clear()
    paragraph.style = "Guide Caption"
    paragraph.add_run(f"Légende proposée : {subject}. À compléter avec les repères visibles de l'écran, les filtres, l'action principale, les tableaux, la pagination et les notifications.")


def add_module(doc, number, module, figure_no):
    doc.add_page_break()
    doc.add_paragraph(f"{number}. {module['title']}", style="Heading 1")
    base.add_label(doc, "Objectif", module["purpose"])
    base.add_label(doc, "Utilisateurs concernés", module["roles"])
    add_text(doc, module["intro"])
    doc.add_paragraph("Résultat attendu", style="Heading 2")
    add_text(doc, "L'utilisateur dispose d'une information fiable et d'une trace exploitable de son action, dans le périmètre de ses habilitations.")
    add_procedure(doc, module["steps"])
    doc.add_paragraph("Points d'attention", style="Heading 2")
    base.add_callout(doc, "Attention", module["attention"], base.ORANGE, "FFF4E6")
    for subject in module["frames"]:
        doc.add_paragraph("Capture à insérer", style="Heading 2")
        add_frame(doc, figure_no, subject)
        figure_no += 1
    return figure_no


def add_roles_table(doc):
    doc.add_paragraph("Profils de référence", style="Heading 2")
    table = doc.add_table(rows=1, cols=2)
    base.set_table_geometry(table, [3000, 6440])
    for cell, text in zip(table.rows[0].cells, ["Profil", "Responsabilité principale"]):
        base.set_cell_shading(cell, base.GREEN)
        base.set_cell_border(cell, base.GREEN_DARK)
        base.set_font(cell.paragraphs[0].add_run(text), 9, base.WHITE, bold=True)
    base.set_repeat_table_header(table.rows[0])
    rows = [
        ("Direction / Présidence", "Pilotage, consultation des indicateurs et décisions selon les délégations."),
        ("Direction Technique", "Supervision métier, validations et coordination des traitements autorisés."),
        ("Point focal BNEC", "Suivi des vérifications, intégrations et qualité des données BNEC."),
        ("Administrateur fonctionnel", "Utilisateurs, référentiels, règles, publications et paramètres fonctionnels autorisés."),
        ("Administrateur système", "Sécurité, continuité, sauvegardes et exploitation technique autorisée."),
        ("Agent de collecte", "Missions, saisie, pièces et soumission des dossiers."),
        ("Agent vérificateur", "Vérification documentaire, confirmations, avis et demandes de complément."),
        ("Cellule de veille", "Relances, suivi, actions et reporting de veille."),
        ("Consultation / audit", "Accès de lecture selon le périmètre accordé."),
    ]
    for profile, responsibility in rows:
        cells = table.add_row().cells
        for cell in cells:
            base.set_cell_border(cell)
        base.set_font(cells[0].paragraphs[0].add_run(profile), 9.5, base.GREEN_DARK, bold=True)
        base.set_font(cells[1].paragraphs[0].add_run(responsibility), 9.5, "222222")


def build():
    doc = Document()
    base.setup_styles(doc)
    section = doc.sections[0]
    base.set_section(section, first_page=True)
    base.add_header_footer(section)

    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_font(p.add_run("HAUQE"), 18, base.GREEN, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_font(p.add_run("SYSTÈME NATIONAL DE GESTION ET DE SUIVI DES CERTIFICATIONS"), 10.5, base.GREEN_DARK, bold=True)
    doc.add_paragraph()
    base.add_title(doc, "Guide global d'utilisation", "SNGSC / HAUQE")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_font(p.add_run("Version rédigée sans captures d'écran"), 11, base.GRAY, italic=True)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_font(p.add_run("Version 1.0 — Août 2026 — Diffusion interne"), 9.5, base.GRAY)
    doc.add_page_break()

    doc.add_paragraph("À propos de ce guide", style="Heading 1")
    add_text(doc, "Ce guide explique l'utilisation fonctionnelle du SNGSC à partir du code local, des feuilles de route et des procédures disponibles au 6 août 2026. Il est destiné aux utilisateurs habilités de la HAUQE.")
    base.add_callout(doc, "Portée", "Les cadres de figure sont volontairement vides : l'équipe HAUQE y insérera des captures nettoyées de toute donnée nominative ou sensible.", base.GREEN, base.GREEN_PALE)
    base.add_callout(doc, "État du système", "Certaines fonctions dépendent encore du raccordement progressif à l'API ou de paramètres métiers publiés. Le guide indique les actions et précautions, sans transformer un paramètre provisoire en règle définitive.", base.ORANGE, "FFF4E6")
    doc.add_paragraph("Fiche documentaire", style="Heading 2")
    base.add_metadata_table(doc)
    doc.add_paragraph("Comment utiliser ce document", style="Heading 2")
    add_text(doc, "Commencez par les chapitres 1 à 4 pour prendre en main le système. Consultez ensuite le module correspondant au statut du dossier. Les écrans et boutons visibles peuvent varier selon les permissions attribuées.")
    add_roles_table(doc)

    doc.add_page_break()
    doc.add_paragraph("Table des matières", style="Heading 1")
    p = doc.add_paragraph()
    base.add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_paragraph("Liste des figures", style="Heading 2")
    p = doc.add_paragraph()
    base.add_field(p, 'TOC \\h \\z \\c "Figure"')
    doc.add_paragraph("Liste des tableaux", style="Heading 2")
    p = doc.add_paragraph()
    base.add_field(p, 'TOC \\h \\z \\c "Tableau"')

    figure_no = 1
    for number, module in enumerate(MODULES, start=1):
        figure_no = add_module(doc, number, module, figure_no)

    doc.add_page_break()
    doc.add_paragraph("Annexes", style="Heading 1")
    doc.add_paragraph("Glossaire essentiel", style="Heading 2")
    glossary = [
        ("BNEC", "Base nationale des entreprises certifiées."),
        ("CVC", "Cellule de Veille des Certifications."),
        ("FUCCS", "Grille de contrôle versionnée utilisée pour documenter le contrôle d'un dossier."),
        ("INFC", "Indice national de fiabilité des certifications, calculé selon un modèle versionné."),
        ("SNCC", "Système national de classement des certificats, associant classe, statut administratif et niveau de risque."),
        ("Traçabilité", "Conservation des actions, décisions, versions, échanges et pièces liées à un dossier."),
    ]
    table = doc.add_table(rows=1, cols=2)
    base.set_table_geometry(table, [2200, 7240])
    for cell, text in zip(table.rows[0].cells, ["Terme", "Définition"]):
        base.set_cell_shading(cell, base.GREEN)
        base.set_cell_border(cell, base.GREEN_DARK)
        base.set_font(cell.paragraphs[0].add_run(text), 9, base.WHITE, bold=True)
    base.set_repeat_table_header(table.rows[0])
    for term, definition in glossary:
        cells = table.add_row().cells
        for cell in cells:
            base.set_cell_border(cell)
        base.set_font(cells[0].paragraphs[0].add_run(term), 9.5, base.GREEN_DARK, bold=True)
        base.set_font(cells[1].paragraphs[0].add_run(definition), 9.5, "222222")
    doc.add_paragraph("Checklist avant action sensible", style="Heading 2")
    checks = [
        "Vérifier le dossier, son statut et les pièces disponibles.",
        "Vérifier que l'action est autorisée par les droits attribués.",
        "Renseigner un motif clair pour toute décision, correction, clôture ou export sensible.",
        "Contrôler les données affichées avant une diffusion ou un export.",
        "Ne jamais inscrire de mot de passe, de code, de jeton ou de secret dans une note ou un document joint.",
    ]
    for text in checks:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
