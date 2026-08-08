from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path("docs_md/guide_global_utilisation_SNGSC_HAUQE_squelette_phase_1.docx")

GREEN = "1F6B45"
GREEN_DARK = "12412A"
GREEN_PALE = "EAF4EE"
ORANGE = "C56A12"
RED = "A52A2A"
GRAY = "6B7280"
LIGHT_GRAY = "F3F4F6"
WHITE = "FFFFFF"
CONTENT_WIDTH = 9360


CHAPTERS = [
    "Page de garde et fiche documentaire",
    "Avertissement, diffusion, historique des versions et lecture du guide",
    "Table des matières, liste des figures, liste des tableaux et glossaire",
    "Présentation du SNGSC et parcours global d'une certification",
    "Connexion, sécurité, MFA, mot de passe oublié et verrouillage de session",
    "Profil, photo, préférences, notifications et actualisation automatique",
    "Tableau de bord, alertes et échéances",
    "Registre national : entreprises, organismes, zones et certifications",
    "Collecte : campagne, mission, déclarant, offres, certifications, preuves et soumission",
    "Vérification documentaire et contrôle FUCCS",
    "Validation N1/N2 et intégration BNEC",
    "Scoring, INFC et classement SNCC",
    "Veille, relances, décisions, actions et communication",
    "Tableaux tactique, stratégique, annuel, national et public",
    "Administration : utilisateurs, référentiels, règles, documents, organismes et publications",
    "Audit, qualité des données, mises à jour BNEC, sauvegarde et restauration",
    "Rapports et exportations",
    "Assistance, erreurs fréquentes et annexes",
]


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color="B7C9BE", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for index, width in enumerate(widths):
        grid.gridCol_lst[index].set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, size=11, color="000000", bold=False, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Mettre à jour dans Word"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Title", 28, GREEN_DARK, 0, 12),
        ("Subtitle", 14, GREEN, 0, 18),
        ("Heading 1", 16, GREEN, 18, 10),
        ("Heading 2", 13, GREEN_DARK, 14, 7),
        ("Heading 3", 11.5, GREEN_DARK, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Subtitle") else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name not in ("Subtitle",)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, color in [("Guide Label", GREEN_DARK), ("Guide Attention", ORANGE), ("Guide Avertissement", RED), ("Guide Caption", GRAY)]:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Aptos"
        style.font.size = Pt(9.5)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name == "Guide Label"
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(4)


def set_section(section, first_page=False):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)
    section.different_first_page_header_footer = first_page


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("SNGSC / HAUQE  |  Guide global d'utilisation")
    set_font(r, 8.5, GRAY, bold=True)
    p.paragraph_format.space_after = Pt(0)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Version 0.1 — Document de travail — Diffusion interne  |  Page ")
    set_font(r, 8.5, GRAY)
    add_page_field(p)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    if subtitle:
        p = doc.add_paragraph(style="Subtitle")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(subtitle)


def add_label(doc, label, text):
    p = doc.add_paragraph(style="Guide Label")
    p.add_run(f"{label} : ")
    p.add_run(text)


def add_callout(doc, heading, text, color, fill):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color)
    p = cell.paragraphs[0]
    r = p.add_run(f"{heading} — ")
    set_font(r, 10, color, bold=True)
    r = p.add_run(text)
    set_font(r, 10, "222222")
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_metadata_table(doc):
    table = doc.add_table(rows=5, cols=2)
    set_table_geometry(table, [2700, 6660])
    labels = ["Titre", "Version", "Statut", "Propriétaire du document", "Dernière mise à jour"]
    values = [
        "Guide global d'utilisation SNGSC / HAUQE",
        "0.1 — Squelette Phase 1",
        "Document modifiable à compléter",
        "HAUQE",
        "6 août 2026",
    ]
    for row, label, value in zip(table.rows, labels, values):
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], GREEN_PALE)
        set_cell_border(row.cells[0])
        set_cell_border(row.cells[1])
        p = row.cells[0].paragraphs[0]
        set_font(p.add_run(label), 9.5, GREEN_DARK, bold=True)
        p = row.cells[1].paragraphs[0]
        set_font(p.add_run(value), 9.5, "222222")


def add_capture_frame(doc, figure_no):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell, GREEN, "12")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(56)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"Figure {figure_no} — Capture à insérer")
    set_font(r, 13, GREEN_DARK, bold=True)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(56)
    r = p.add_run("Masquer les données nominatives, identifiants, adresses électroniques et informations sensibles avant insertion.")
    set_font(r, 9, GRAY, italic=True)
    p = doc.add_paragraph(style="Guide Caption")
    p.add_run("Légende à compléter : titre ; filtres ; bouton principal ; tableau ou cartes ; actions ; pagination ; notifications.")


def add_procedure_template(doc):
    doc.add_paragraph("Procédure pas à pas", style="Heading 2")
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [900, 3000, 5460])
    headers = ["Étape", "Action de l'utilisateur", "Résultat attendu"]
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, GREEN)
        set_cell_border(cell, GREEN_DARK)
        p = cell.paragraphs[0]
        set_font(p.add_run(label), 9, WHITE, bold=True)
    set_repeat_table_header(table.rows[0])
    for number in range(1, 4):
        cells = table.add_row().cells
        for cell in cells:
            set_cell_border(cell)
        set_font(cells[0].paragraphs[0].add_run(str(number)), 9.5, GREEN_DARK, bold=True)
        set_font(cells[1].paragraphs[0].add_run("À compléter"), 9.5, GRAY, italic=True)
        set_font(cells[2].paragraphs[0].add_run("À compléter"), 9.5, GRAY, italic=True)


def add_module_template(doc, chapter_no, title, figure_no):
    doc.add_paragraph(f"{chapter_no}. {title}", style="Heading 1")
    add_label(doc, "Objectif", "À compléter.")
    add_label(doc, "Utilisateurs concernés", "À compléter.")
    add_label(doc, "Prérequis", "À compléter.")
    doc.add_paragraph("Présentation de l'interface", style="Heading 2")
    doc.add_paragraph("Décrire brièvement la page, les zones de travail et les actions disponibles.")
    add_capture_frame(doc, figure_no)
    doc.add_paragraph("Légende numérotée", style="Heading 2")
    doc.add_paragraph("1. À compléter.\n2. À compléter.\n3. À compléter.")
    add_procedure_template(doc)
    doc.add_paragraph("Résultat attendu", style="Heading 2")
    doc.add_paragraph("À compléter.")
    add_callout(doc, "Point d'attention", "À compléter : contrôles, limites ou erreurs fréquentes.", ORANGE, "FFF4E6")
    doc.add_paragraph("Liens avec les autres modules", style="Heading 2")
    doc.add_paragraph("À compléter.")


def build():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    set_section(section, first_page=True)
    add_header_footer(section)

    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HAUQE")
    set_font(r, 17, GREEN, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SYSTÈME NATIONAL DE GESTION ET DE SUIVI DES CERTIFICATIONS")
    set_font(r, 10.5, GREEN_DARK, bold=True)
    doc.add_paragraph()
    add_title(doc, "Guide global d'utilisation", "SNGSC / HAUQE")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Squelette professionnel et modifiable — Phase 1")
    set_font(r, 11, GRAY, italic=True)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Version 0.1 | Août 2026 | Diffusion interne")
    set_font(r, 9.5, GRAY)
    doc.add_page_break()

    doc.add_paragraph("Fiche documentaire", style="Heading 1")
    add_metadata_table(doc)
    add_callout(doc, "Usage", "Ce guide est un modèle de travail. Les contenus fonctionnels et les captures sont ajoutés progressivement par la HAUQE.", GREEN, GREEN_PALE)

    doc.add_paragraph("Avertissement et diffusion", style="Heading 1")
    add_callout(doc, "Confidentialité", "Ne pas intégrer de secrets, mots de passe, données personnelles ni détails techniques sensibles.", RED, "FDECEC")
    add_label(doc, "Diffusion", "Interne — agents habilités de la HAUQE.")
    doc.add_paragraph("Historique des versions", style="Heading 2")
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1500, 2000, 5860])
    for cell, label in zip(table.rows[0].cells, ["Version", "Date", "Évolution"]):
        set_cell_shading(cell, GREEN)
        set_font(cell.paragraphs[0].add_run(label), 9, WHITE, bold=True)
        set_cell_border(cell, GREEN_DARK)
    set_repeat_table_header(table.rows[0])
    cells = table.add_row().cells
    for cell in cells: set_cell_border(cell)
    for cell, value in zip(cells, ["0.1", "06/08/2026", "Création du squelette Phase 1"]):
        set_font(cell.paragraphs[0].add_run(value), 9.5, "222222")
    doc.add_paragraph("Comment lire ce guide", style="Heading 2")
    doc.add_paragraph("Chaque module suit la même structure : objectif, utilisateurs concernés, prérequis, interface, capture, procédure, résultat et points d'attention.")

    doc.add_paragraph("Navigation du document", style="Heading 1")
    doc.add_paragraph("Table des matières", style="Heading 2")
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_paragraph("Liste des figures", style="Heading 2")
    p = doc.add_paragraph()
    add_field(p, 'TOC \\h \\z \\c "Figure"')
    doc.add_paragraph("Liste des tableaux", style="Heading 2")
    p = doc.add_paragraph()
    add_field(p, 'TOC \\h \\z \\c "Tableau"')
    doc.add_paragraph("Glossaire", style="Heading 2")
    doc.add_paragraph("À compléter au fil de la rédaction fonctionnelle.")

    figure_no = 1
    for index, chapter in enumerate(CHAPTERS[3:], start=4):
        doc.add_page_break()
        add_module_template(doc, index, chapter, figure_no)
        figure_no += 1

    doc.add_page_break()
    doc.add_paragraph("Annexes", style="Heading 1")
    doc.add_paragraph("Annexe A — Glossaire détaillé", style="Heading 2")
    doc.add_paragraph("À compléter.")
    doc.add_paragraph("Annexe B — Références et contacts d'assistance", style="Heading 2")
    doc.add_paragraph("À compléter.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
