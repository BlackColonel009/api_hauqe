from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("output/docx")
OUT.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill="E7E7E7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="B5B5B5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def font_black(run, bold=False, size=None, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)


def base_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)

    for name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
    title_pr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_pr.remove(title_border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("HAUQE Certif - Support de formation | Page ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_black(p.add_run(title), bold=True, size=19)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_black(p.add_run(subtitle), size=10.5, italic=True)


def add_info_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(12.3)
    for label, value in rows:
        cells = table.add_row().cells
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
        set_cell_shading(cells[0], "E7E7E7")
        font_black(cells[0].paragraphs[0].add_run(label), bold=True, size=9.5)
        font_black(cells[1].paragraphs[0].add_run(value), size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(doc, title, entries):
    p = doc.add_paragraph()
    font_black(p.add_run(title), bold=True, size=10.5)
    for entry in entries:
        item = doc.add_paragraph(style="List Bullet")
        item.paragraph_format.space_after = Pt(2)
        font_black(item.add_run(entry), size=9.7)


def add_notes(doc, label, lines=2):
    p = doc.add_paragraph()
    font_black(p.add_run(label), bold=True, size=9.5)
    for _ in range(lines):
        p = doc.add_paragraph("________________________________________________________________________________")
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            font_black(run, size=8.5)


CASES = [
    {
        "title": "Création complète d une entreprise de production",
        "certificate": "FC 01 - AgroNoura SARL",
        "scenario": "La société fictive AgroNoura transforme du manioc à Kpalimé. L agent doit créer la fiche, saisir l entreprise, une offre et une certification déclarée.",
        "tasks": ["Créer ou sélectionner l entreprise AgroNoura SARL.", "Renseigner la mission, la zone administrative et la priorité normale.", "Ajouter l offre Farine de manioc et la certification ISO 22000 déclarée.", "Enregistrer le brouillon puis vérifier la complétude avant soumission."],
        "expected": "La fiche contient une entreprise unique, une offre, une certification et une mission cohérente.",
        "note": "Contrôler que les champs d identité, mission et certification sont renseignés avant la soumission.",
    },
    {
        "title": "Entreprise sans RCCM NIF ni IFU",
        "certificate": "FC 02 - Savane Miel Coopérative",
        "scenario": "La coopérative fictive Savane Miel ne communique pas encore ses références légales. Les champs RCCM NIF et IFU sont désactivés par l administration.",
        "tasks": ["Créer l entreprise avec son nom, sa zone et son contact disponible.", "Ne pas inventer de RCCM, NIF ou IFU.", "Ajouter l observation Références légales non communiquées lors de la collecte.", "Enregistrer la fiche en brouillon."],
        "expected": "La collecte reste exploitable sans blocage et sans données légales fictives.",
        "note": "Rappeler que ces champs sont facultatifs selon le paramètre HAUQE et peuvent être activés plus tard.",
    },
    {
        "title": "Réutilisation d une entreprise déjà enregistrée",
        "certificate": "FC 03 - TogoCacao Industrie",
        "scenario": "TogoCacao Industrie apparaît dans les résultats de recherche avec le même numéro de téléphone et la même localité.",
        "tasks": ["Rechercher TogoCacao Industrie avant toute précréation.", "Sélectionner l entreprise existante.", "Vérifier son identifiant HAUQE et ses coordonnées connues.", "Ajouter uniquement les nouvelles déclarations liées à la mission."],
        "expected": "Aucun doublon entreprise n est créé.",
        "note": "Faire verbaliser la différence entre sélectionner une entreprise et créer une nouvelle entreprise.",
    },
    {
        "title": "Campagne sélectionnée avec informations partielles",
        "certificate": "FC 09 - Formation Atlas (utiliser uniquement les données entreprise et mission ; les dates de campagne sont fournies par le formateur)",
        "scenario": "Une campagne existante est choisie. Elle possède un code, un objet et une date de début mais pas de date de fin.",
        "tasks": ["Sélectionner la campagne existante.", "Vérifier le préremplissage de la référence de mission, de l objet et de la date de début.", "Laisser la date de fin vide si elle n a pas été communiquée.", "Compléter la zone administrative et la priorité si elles sont connues."],
        "expected": "Les données de campagne sont reprises sans invention de date manquante.",
        "note": "Insister sur la relecture des données préremplies avant la suite de la saisie.",
    },
    {
        "title": "Ajout de plusieurs offres déclarées",
        "certificate": "FC 04 - Kara Textile",
        "scenario": "Kara Textile produit des uniformes scolaires, des vêtements de travail et des sacs en tissu.",
        "tasks": ["Ajouter une offre par rubrique.", "Utiliser Ajouter une offre pour créer trois lignes distinctes.", "Renseigner la désignation et toute capacité connue.", "Supprimer une ligne vide créée par erreur."],
        "expected": "Les trois offres restent séparées et aucune ligne vide ne subsiste.",
        "note": "Demander au participant de vérifier le fonctionnement de l icône de suppression avant enregistrement.",
    },
    {
        "title": "Deux certifications pour une même entreprise",
        "certificate": "FC 05 - Clinique Horizon ; compléter par une seconde déclaration ISO 9001 donnée oralement par le formateur",
        "scenario": "Clinique Horizon déclare ISO 9001 et ISO 15189 avec des portées différentes.",
        "tasks": ["Ajouter une certification déclarée par ligne.", "Renseigner les noms, numéros, normes et portées quand elles sont disponibles.", "Saisir les dates propres à chaque certificat.", "Contrôler qu aucune information ISO 9001 ne se retrouve sur ISO 15189."],
        "expected": "Deux déclarations indépendantes sont enregistrées pour la même entreprise.",
        "note": "Faire repérer les champs qui appartiennent à la certification et non à l entreprise.",
    },
    {
        "title": "Certificat sans date d expiration",
        "certificate": "FC 06 - BatiPlus Togo",
        "scenario": "BatiPlus Togo fournit un certificat mais la date de fin est absente du document remis.",
        "tasks": ["Saisir les informations visibles du certificat.", "Laisser la date de fin vide.", "Joindre le document si disponible.", "Ajouter une observation Date d expiration à confirmer auprès de l organisme."],
        "expected": "La fiche est exacte et prépare une vérification documentaire ultérieure.",
        "note": "Une date inconnue ne doit jamais être estimée ou remplacée par la date du jour.",
    },
    {
        "title": "Certificat expiré",
        "certificate": "FC 07 - Laiterie du Plateau",
        "scenario": "Laiterie du Plateau présente un certificat dont la date de fin est dépassée depuis deux mois.",
        "tasks": ["Saisir les dates telles qu elles figurent sur le certificat.", "Ajouter le statut ou observation Certificat expiré selon le formulaire.", "Mettre une priorité élevée si la règle métier le justifie.", "Soumettre avec une note destinée à la vérification."],
        "expected": "Le dossier signale la situation sans effacer la déclaration de l entreprise.",
        "note": "Expliquer que la collecte constate ; la vérification et la décision suivent dans leurs modules respectifs.",
    },
    {
        "title": "Écart entre raison sociale et certificat",
        "certificate": "FC 08 - TechNova Services",
        "scenario": "L entreprise est enregistrée comme TechNova Services, mais le certificat porte Tech Nova Service SARL.",
        "tasks": ["Conserver la raison sociale déjà connue dans l entreprise.", "Saisir le libellé figurant sur le certificat dans la déclaration ou l observation.", "Ne pas créer une seconde entreprise.", "Signaler l écart de dénomination pour vérification."],
        "expected": "Le rapprochement est possible et l écart est tracé.",
        "note": "Faire distinguer une variante de nom d un nouveau sujet juridique.",
    },
    {
        "title": "Détection d une certification possiblement en doublon",
        "certificate": "FC 12 - Logistique Delta",
        "scenario": "Le numéro LD QMS 2026 004 apparaît déjà dans une autre fiche, mais le participant ne sait pas si le certificat a été renouvelé.",
        "tasks": ["Rechercher le numéro de certification et l entreprise.", "Comparer norme, organisme, dates et portée.", "Ne pas saisir une nouvelle ligne identique avant vérification.", "Ajouter une observation Possible renouvellement ou doublon à examiner."],
        "expected": "Le dossier évite un doublon tout en conservant la question à résoudre.",
        "note": "Le doublon est une alerte de contrôle, pas une preuve automatique d erreur.",
    },
    {
        "title": "Mission prioritaire dans une zone différente",
        "certificate": "FC 02 - Savane Miel Coopérative (adapter la zone et la priorité selon la consigne du cas)",
        "scenario": "Une mission de collecte doit être préparée pour une entreprise de la région des Savanes avec priorité haute.",
        "tasks": ["Choisir la zone administrative correcte.", "Renseigner la priorité haute et son motif dans les observations.", "Vérifier la cohérence entre campagne, mission et zone.", "Enregistrer avant de poursuivre la saisie détaillée."],
        "expected": "La mission est localisable et priorisée de manière justifiée.",
        "note": "La priorité doit refléter une situation réelle : échéance, risque, campagne ciblée ou demande HAUQE.",
    },
    {
        "title": "Document illisible ou incomplet",
        "certificate": "FC 13 - Pharma Solaire (le formateur fournit une copie volontairement floue ou cache le numéro avant distribution)",
        "scenario": "Le scan d un certificat est flou et son numéro ne peut pas être lu.",
        "tasks": ["Joindre le document lorsqu il est autorisé.", "Saisir uniquement les informations lisibles.", "Laisser le numéro vide s il est illisible.", "Ajouter une observation Document illisible, original ou copie nette à demander."],
        "expected": "Le dossier reste honnête et oriente la vérification vers la bonne action.",
        "note": "Ne jamais deviner un numéro de certificat à partir d une image incertaine.",
    },
    {
        "title": "Date de début postérieure à la date de fin",
        "certificate": "FC 14 - Grains de Paix (le formateur dicte la saisie erronée afin de faire repérer l incohérence)",
        "scenario": "Le participant saisit par erreur 15 novembre 2026 comme début et 15 novembre 2025 comme fin.",
        "tasks": ["Repérer l incohérence avant l enregistrement.", "Corriger les dates à partir du document source.", "Si le document est ambigu, laisser la donnée litigieuse vide et noter le besoin de confirmation.", "Recontrôler la période affichée."],
        "expected": "Aucune période impossible ne doit être validée.",
        "note": "Utiliser ce cas pour montrer le rôle du contrôle de cohérence et de la relecture humaine.",
    },
    {
        "title": "Organisme certificateur renseigné organisme accréditeur absent",
        "certificate": "FC 13 - Pharma Solaire",
        "scenario": "Le certificat cite Santé Certification, mais aucune information ne permet d identifier son accréditeur.",
        "tasks": ["Sélectionner ou créer uniquement l organisme certificateur connu.", "Laisser l accréditeur non renseigné.", "Noter Accréditation à confirmer dans les observations.", "Préparer le dossier pour une confirmation externe si nécessaire."],
        "expected": "Les rôles certificateur et accréditeur ne sont pas confondus.",
        "note": "Le certificateur délivre ; l accréditeur reconnaît la compétence du certificateur selon son périmètre.",
    },
    {
        "title": "Portée de certification trop générale",
        "certificate": "FC 11 - Jus du Togo (le formateur masque la portée détaillée et communique seulement Activités agroalimentaires)",
        "scenario": "Le certificat mentionne seulement Activités agroalimentaires alors que l entreprise produit du jus et des confitures.",
        "tasks": ["Saisir la portée exactement telle qu elle apparaît.", "Ajouter les produits dans les offres déclarées.", "Ne pas élargir la portée par interprétation.", "Signaler le besoin de précision si la portée doit être vérifiée."],
        "expected": "Les produits et la portée restent deux informations distinctes.",
        "note": "Une portée vague peut demander une vérification documentaire, sans rendre le certificat automatiquement invalide.",
    },
    {
        "title": "Entreprise sans certification déclarée",
        "certificate": "FC 10 - Menuiserie Moderne",
        "scenario": "Menuiserie Moderne accepte la collecte mais indique ne détenir aucune certification.",
        "tasks": ["Créer ou sélectionner l entreprise.", "Renseigner mission, offre et contact.", "Ne pas créer de certification fictive.", "Utiliser l observation Aucune certification déclarée au moment de la collecte."],
        "expected": "La fiche est valable même sans certification déclarée.",
        "note": "Le système recense aussi les entreprises à vérifier ou à accompagner.",
    },
    {
        "title": "Brouillon à réinitialiser",
        "certificate": "FC 04 - Kara Textile (démarrer la fiche avec une mauvaise entreprise ou offre, puis la réinitialiser)",
        "scenario": "Un agent a commencé une fiche sur la mauvaise entreprise et aucune donnée ne doit être conservée.",
        "tasks": ["Vérifier que la fiche est encore en brouillon.", "Utiliser Réinitialiser le brouillon dans le formulaire.", "Contrôler que campagne, mission, entreprise, offres et certifications peuvent être ressaisies.", "Ne pas utiliser la réinitialisation sur une fiche déjà soumise."],
        "expected": "Les entrées du formulaire sont nettoyées pour recommencer la collecte.",
        "note": "La réinitialisation concerne la saisie courante ; elle ne supprime pas l historique métier validé.",
    },
    {
        "title": "Identifiant temporaire et code définitif BNEC",
        "certificate": "FC 01 - AgroNoura SARL",
        "scenario": "Une entreprise précréée pendant la collecte reçoit un identifiant provisoire HAUQE ENT année séquence.",
        "tasks": ["Constater l identifiant proposé sans le modifier arbitrairement.", "Terminer la collecte et soumettre la fiche.", "Expliquer que le code institutionnel définitif est attribué à l intégration BNEC.", "Vérifier ensuite la traçabilité entre l identifiant provisoire et le code final."],
        "expected": "Le participant comprend la différence entre identification de collecte et codification officielle BNEC.",
        "note": "Ne jamais utiliser un ancien format TEMP COL pour une entreprise nouvellement précréée.",
    },
    {
        "title": "Révision après retour de vérification",
        "certificate": "FC 08 - TechNova Services",
        "scenario": "La fiche soumise revient avec une demande de correction sur le numéro de certificat et la portée.",
        "tasks": ["Ouvrir la fiche concernée et créer une nouvelle révision si le statut le permet.", "Corriger seulement les informations demandées à partir d une preuve fiable.", "Conserver le motif de révision.", "Resoumettre après contrôle de complétude."],
        "expected": "La correction est tracée sans effacer la version initiale.",
        "note": "Faire rappeler le parcours : une correction peut nécessiter une nouvelle vérification et les étapes aval sont réévaluées.",
    },
    {
        "title": "Fiche complète prête à soumettre",
        "certificate": "FC 09 - Formation Atlas",
        "scenario": "Le dossier de Formation Atlas contient entreprise, mission, campagne, zone, deux offres, une certification et une pièce jointe lisible.",
        "tasks": ["Revoir tous les champs obligatoires.", "Contrôler les dates et l organisme certificateur.", "Vérifier qu il n existe pas de ligne offre ou certification vide.", "Soumettre la fiche et lire le statut obtenu."],
        "expected": "Le dossier passe de brouillon à soumise et devient disponible pour la vérification documentaire.",
        "note": "Clore la mise en pratique en demandant au participant quelle est la prochaine étape et quel rôle la réalise.",
    },
]


CERTIFICATES = [
    ("FC 01", "AgroNoura SARL", "HAUQE-ENT-2026-0001", "ISO 22000 2018", "QMS-AN-2026-014", "Transformation de manioc et production de farine alimentaire", "01/02/2026", "31/01/2029", "Bureau Certif Afrique", "Accredia Togo", "Kpalimé", "Copie lisible"),
    ("FC 02", "Savane Miel Coopérative", "HAUQE-ENT-2026-0002", "HACCP", "HACCP-SM-118", "Récolte, conditionnement et commercialisation de miel", "15/03/2026", "14/03/2027", "QualiCert Ouest", "Accréditation non communiquée", "Dapaong", "RCCM NIF IFU non communiqués"),
    ("FC 03", "TogoCacao Industrie", "HAUQE-ENT-2025-0048", "ISO 9001 2015", "QMS-TCI-2025-077", "Transformation et conditionnement de cacao", "10/06/2025", "09/06/2028", "CertiAfrique", "COFRAC Fiction", "Lomé", "Entreprise déjà enregistrée"),
    ("FC 04", "Kara Textile", "HAUQE-ENT-2026-0004", "ISO 9001 2015", "KT-Q-2026-032", "Confection d uniformes, vêtements de travail et sacs textiles", "01/05/2026", "30/04/2029", "Bureau Certif Afrique", "Accredia Togo", "Kara", "Trois offres à enregistrer"),
    ("FC 05", "Clinique Horizon", "HAUQE-ENT-2026-0005", "ISO 15189 2022", "LAB-CH-2026-091", "Examens de biologie médicale et analyses de laboratoire", "20/01/2026", "19/01/2028", "Santé Certification", "TogoAccred", "Lomé", "Deuxième certificat ISO 9001 à saisir séparément"),
    ("FC 06", "BatiPlus Togo", "HAUQE-ENT-2026-0006", "ISO 45001 2018", "OHS-BPT-443", "Construction de bâtiments et travaux publics", "05/07/2025", "Non communiquée", "CertiBat Afrique", "Accredia BTP", "Sokodé", "Date de fin absente"),
    ("FC 07", "Laiterie du Plateau", "HAUQE-ENT-2026-0007", "ISO 22000 2018", "FS-LP-2023-025", "Production de lait pasteurisé et yaourt", "01/04/2023", "31/03/2026", "QualiCert Ouest", "TogoAccred", "Atakpamé", "Certificat expiré"),
    ("FC 08", "TechNova Services", "HAUQE-ENT-2026-0008", "ISO 9001 2015", "TN-QMS-26-18", "Développement logiciel et maintenance informatique", "18/08/2026", "17/08/2029", "Digital Quality Board", "Accredia Numérique", "Lomé", "Certificat au nom Tech Nova Service SARL"),
    ("FC 09", "Formation Atlas", "HAUQE-ENT-2026-0009", "ISO 21001 2018", "EDU-FA-2026-010", "Formation professionnelle continue et accompagnement pédagogique", "01/09/2026", "31/08/2029", "EduCert Afrique", "Accredia Education", "Lomé", "Dossier complet"),
    ("FC 10", "Menuiserie Moderne", "HAUQE-ENT-2026-0010", "Aucune certification déclarée", "Non applicable", "Fabrication de meubles sur mesure", "Non applicable", "Non applicable", "Non applicable", "Non applicable", "Tsévié", "Cas sans certification"),
    ("FC 11", "Jus du Togo", "HAUQE-ENT-2026-0011", "FSSC 22000", "FSSC-JT-2026-059", "Production de jus de fruits et confitures", "12/02/2026", "11/02/2029", "Food Cert International", "Accredia Agro", "Lomé", "Portée à comparer avec les offres"),
    ("FC 12", "Logistique Delta", "HAUQE-ENT-2026-0012", "ISO 9001 2015", "LD-QMS-2026-004", "Transport routier, entreposage et distribution", "01/01/2026", "31/12/2028", "CertiAfrique", "COFRAC Fiction", "Lomé", "Numéro à rechercher pour doublon"),
    ("FC 13", "Pharma Solaire", "HAUQE-ENT-2026-0013", "ISO 13485 2016", "MED-PS-2026-220", "Distribution de dispositifs médicaux", "03/06/2026", "02/06/2029", "Santé Certification", "Accréditation non communiquée", "Lomé", "Accréditeur manquant"),
    ("FC 14", "Grains de Paix", "HAUQE-ENT-2026-0014", "Agriculture biologique", "BIO-GP-2026-17", "Production et conditionnement de céréales biologiques", "Non communiquée", "30/09/2027", "Bio Cert Sahel", "West Africa Organic Accreditation", "Mango", "Date de début absente"),
    ("FC 15", "EcoPalm Industries", "HAUQE-ENT-2026-0015", "ISO 14001 2015", "EMS-EPI-2026-005", "Extraction d huile de palme et gestion environnementale du site", "01/03/2026", "28/02/2029", "GreenCert Afrique", "Accredia Environnement", "Kpalimé", "Dossier complet environnement"),
    ("FC 16", "AquaPure Togo", "HAUQE-ENT-2026-0016", "ISO 22000 2018", "FS-APT-2026-088", "Production et conditionnement d eau potable", "10/04/2026", "09/04/2029", "Food Cert International", "Accredia Agro", "Lomé", "Dates et portée à relever"),
    ("FC 17", "Boulangerie Zio", "HAUQE-ENT-2026-0017", "HACCP", "Non communiqué", "Fabrication de pain, viennoiseries et pâtisseries", "12/01/2026", "11/01/2027", "QualiCert Ouest", "TogoAccred", "Tsévié", "Numéro du certificat absent"),
    ("FC 18", "PlastTogo Recyclage", "HAUQE-ENT-2026-0018", "ISO 9001 2015", "QMS-PTR-2026-101", "Collecte, recyclage et transformation de matières plastiques", "05/05/2026", "04/05/2029", "CertiAfrique", "Accréditation non communiquée", "Lomé", "Accréditeur à confirmer"),
    ("FC 19", "Ferme Niamtougou", "HAUQE-ENT-2026-0019", "GLOBALG.A.P. IFA", "GAP-FN-2026-040", "Production de légumes frais et cultures maraîchères", "01/06/2026", "31/05/2027", "Agro Certification Sahel", "West Africa Organic Accreditation", "Niamtougou", "Certification agricole valide"),
    ("FC 20", "Transit Glo", "HAUQE-ENT-2026-0020", "ISO 28000 2022", "SEC-TG-2026-073", "Gestion de la sécurité de la chaîne logistique", "15/02/2026", "14/02/2029", "Secure Chain Certification", "Accredia Logistique", "Lomé", "Portée logistique à comparer"),
    ("FC 21", "Imprimerie Kanta", "HAUQE-ENT-2026-0021", "ISO 9001 2015", "QMS-IK-2025-219", "Impression numérique, offset et façonnage", "20/10/2025", "19/10/2028", "Bureau Certif Afrique", "COFRAC Fiction", "Kara", "Entreprise et certificat cohérents"),
    ("FC 22", "Hôtel Koudjou", "HAUQE-ENT-2026-0022", "ISO 21401 2018", "HOS-HK-2026-012", "Hébergement touristique, restauration et services de conférence", "Non communiquée", "30/06/2029", "Tourism Quality Africa", "Accredia Services", "Lomé", "Date de début absente"),
    ("FC 23", "BioPharm Lomé", "HAUQE-ENT-2026-0023", "Bonnes pratiques de fabrication", "BPF-BPL-2026-062", "Fabrication et conditionnement de compléments alimentaires", "01/08/2026", "Non communiquée", "Santé Certification", "TogoAccred", "Lomé", "Date de fin à confirmer"),
    ("FC 24", "Ciment du Mono", "HAUQE-ENT-2026-0024", "ISO 45001 2018", "OHS-CDM-2026-145", "Production de ciment et exploitation de carrière", "07/07/2026", "06/07/2029", "CertiBat Afrique", "Accredia BTP", "Tabligbo", "Santé et sécurité au travail"),
    ("FC 25", "Kpalimé Fruits", "HAUQE-ENT-2026-0025", "FSSC 22000", "FSSC-KF-2026-084", "Transformation de fruits tropicaux et conditionnement de jus", "18/03/2026", "17/03/2029", "Food Cert International", "Accredia Agro", "Kpalimé", "Produits et portée à rapprocher"),
    ("FC 26", "DataLink Togo", "HAUQE-ENT-2026-0026", "ISO IEC 27001 2022", "ISMS-DLT-2026-031", "Hébergement de données, développement logiciel et support informatique", "01/09/2026", "31/08/2029", "Digital Quality Board", "Accredia Numérique", "Lomé", "Certification de sécurité de l information"),
    ("FC 27", "Coopérative Avenir", "HAUQE-ENT-2026-0027", "Fairtrade Standard", "FT-CA-2024-056", "Production et commercialisation de cacao équitable", "01/01/2024", "31/12/2025", "FairTrade Cert Afrique", "Accréditation non communiquée", "Atakpamé", "Certificat expiré"),
    ("FC 28", "Soja Plus", "HAUQE-ENT-2026-0028", "Agriculture biologique", "BIO-SP-2026-009", "Production de soja biologique et transformation en farine", "22/05/2026", "21/05/2027", "Bio Cert Sahel", "West Africa Organic Accreditation", "Sokodé", "Certification biologique valide"),
    ("FC 29", "Atelier Meka", "HAUQE-ENT-2026-0029", "Aucune certification déclarée", "Non applicable", "Fabrication de mobilier métallique et menuiserie aluminium", "Non applicable", "Non applicable", "Non applicable", "Non applicable", "Lomé", "Cas sans certification déclarée"),
    ("FC 30", "Togo Céramique", "HAUQE-ENT-2026-0030", "ISO 9001 2015", "QMS-TC-2026-116", "Fabrication de carreaux et produits céramiques", "11/11/2026", "10/11/2029", "Bureau Certif Afrique", "COFRAC Fiction", "Kara", "Vérifier les dates et l organisme certificateur"),
]


def create_cases_document():
    doc = base_document()
    add_title(doc, "Cas pratiques de saisie de fiche de collecte", "SNGSC HAUQE Certif - Formation des utilisateurs")
    p = doc.add_paragraph()
    font_black(p.add_run("Utilisation du support. "), bold=True, size=10.2)
    font_black(p.add_run("Chaque cas est fictif. Le formateur attribue un cas au participant ou à un groupe, puis utilise la note de contrôle pour la restitution. Les données incomplètes, les anomalies et les incohérences doivent être traitées sans invention de données."), size=10.2)
    add_info_table(doc, [
        ("Durée indicative", "8 à 12 minutes par cas, puis restitution collective."),
        ("Objectif", "Saisir une fiche fiable, reconnaître les anomalies et identifier la prochaine étape du dossier."),
        ("Règle de formation", "Toute donnée absente, illisible ou incertaine est laissée vide et expliquée dans les observations."),
    ])
    doc.add_page_break()

    for index, case in enumerate(CASES, 1):
        doc.add_paragraph(f"Cas pratique {index:02d}", style="Heading 1")
        p = doc.add_paragraph()
        font_black(p.add_run(case["title"]), bold=True, size=11)
        p = doc.add_paragraph()
        font_black(p.add_run("Situation : "), bold=True, size=9.8)
        font_black(p.add_run(case["scenario"]), size=9.8)
        p = doc.add_paragraph()
        font_black(p.add_run("Document à utiliser : "), bold=True, size=9.8)
        font_black(p.add_run(case["certificate"]), size=9.8)
        add_bullets(doc, "Travail demandé", case["tasks"])
        p = doc.add_paragraph()
        font_black(p.add_run("Résultat attendu : "), bold=True, size=9.7)
        font_black(p.add_run(case["expected"]), size=9.7)
        p = doc.add_paragraph()
        font_black(p.add_run("Note de contrôle du formateur : "), bold=True, size=9.7)
        font_black(p.add_run(case["note"]), size=9.7)
        add_notes(doc, "Notes du participant", 4)
        if index != len(CASES):
            doc.add_page_break()

    doc.save(OUT / "Cas_pratiques_saisie_fiche_collecte_SNGSC_HAUQE.docx")


def create_certificates_document():
    doc = base_document()
    for index, cert in enumerate(CERTIFICATES):
        ref, company, identifier, standard, number, scope, start, end, certifier, accreditor, location, note = cert
        if index:
            doc.add_page_break()
        add_title(doc, "Certificat fictif pour exercice de formation", "SNGSC HAUQE Certif - Document sans valeur juridique")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_black(p.add_run(ref), bold=True, size=12)
        doc.add_paragraph()
        add_info_table(doc, [
            ("Entreprise", company),
            ("Identifiant HAUQE", identifier),
            ("Localité", location),
            ("Norme ou référentiel", standard),
            ("Numéro du certificat", number),
            ("Portée", scope),
            ("Date de début", start),
            ("Date de fin", end),
            ("Organisme certificateur", certifier),
            ("Organisme accréditeur", accreditor),
        ])
        doc.add_paragraph()
        p = doc.add_paragraph()
        font_black(p.add_run("Indication pour l exercice : "), bold=True, size=10)
        font_black(p.add_run(note), size=10)
        add_notes(doc, "Informations à saisir ou à signaler", 4)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_black(p.add_run("DOCUMENT FICTIF DE FORMATION - NE PAS UTILISER COMME CERTIFICAT OFFICIEL"), bold=True, size=8.5)
    doc.save(OUT / "Certificats_fictifs_formation_SNGSC_HAUQE.docx")


if __name__ == "__main__":
    create_cases_document()
    create_certificates_document()
